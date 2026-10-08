"""Pruebas de la lógica contable y de las páginas. Ejecutar con:
    .venv\\Scripts\\python -m unittest discover tests
"""
import io
import sqlite3
import tempfile
import unittest
from pathlib import Path

from contabilidad import create_app, db, importar_creandogestion, nucleo
from contabilidad.nucleo import ErrorValidacion


class BaseConDatos(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.ruta = Path(self.tmp.name) / "prueba.db"
        self.app = create_app(self.ruta)
        self.conn = db.conectar(self.ruta)
        self.emp = self.conn.execute("SELECT MIN(id) FROM empresas").fetchone()[0]

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def id_cuenta(self, codigo, empresa=None):
        return self.conn.execute("SELECT id FROM cuentas WHERE empresa_id = ? AND codigo = ?",
                                 (empresa or self.emp, codigo)).fetchone()[0]

    def asiento(self, tipo, fecha, glosa, *movs, empresa=None, apertura=False):
        """movs: (codigo, debe, haber)"""
        e = empresa or self.emp
        lineas = [{"cuenta_id": self.id_cuenta(c, e), "glosa": "", "debe": d, "haber": h} for c, d, h in movs]
        return nucleo.guardar_asiento(self.conn, e, tipo, fecha, glosa, lineas, apertura=apertura)

    def cargar_ejemplo(self, empresa=None):
        e = {"empresa": empresa}
        # Apertura: aporte de capital al banco
        self.asiento("TRASPASO", "2026-01-01", "Apertura", ("1.1.02", 10_000_000, 0), ("3.1.01", 0, 10_000_000), **e)
        # Venta con IVA 19 %
        self.asiento("INGRESO", "2026-02-10", "Venta factura 1",
                     ("1.1.02", 1_190_000, 0), ("4.1.01", 0, 1_000_000), ("2.1.04", 0, 190_000), **e)
        self.asiento("EGRESO", "2026-02-15", "Arriendo febrero", ("5.2.04", 400_000, 0), ("1.1.02", 0, 400_000), **e)
        self.asiento("EGRESO", "2026-03-01", "Compra computador", ("1.2.06", 800_000, 0), ("1.1.09", 152_000, 0),
                     ("1.1.02", 0, 952_000), **e)

    def nueva_empresa(self, nombre="Otra SpA", rut="11.111.111-1"):
        return nucleo.crear_empresa(self.conn, {"razon_social": nombre, "rut": rut, "ejercicio": "2026"})


class PruebasFormatos(unittest.TestCase):
    def test_rut(self):
        self.assertTrue(nucleo.validar_rut("12.345.678-5"))
        self.assertTrue(nucleo.validar_rut("76086428-5"))
        self.assertTrue(nucleo.validar_rut("11.111.111-1"))
        self.assertFalse(nucleo.validar_rut("12.345.678-9"))
        self.assertFalse(nucleo.validar_rut("abc"))
        self.assertEqual(nucleo.formatear_rut("123456785"), "12.345.678-5")

    def test_montos(self):
        self.assertEqual(nucleo.parse_monto("1.234.567"), 1234567)
        self.assertEqual(nucleo.parse_monto("$ 1234567"), 1234567)
        self.assertEqual(nucleo.parse_monto(""), 0)
        for malo in ("12,5", "1.5", "-100", "abc", "1.23.456"):
            with self.assertRaises(ErrorValidacion, msg=malo):
                nucleo.parse_monto(malo)
        self.assertEqual(nucleo.formato_clp(1234567), "1.234.567")
        self.assertEqual(nucleo.formato_clp(-5000), "-5.000")


class PruebasAsientos(BaseConDatos):
    def test_rechaza_descuadrado(self):
        with self.assertRaises(ErrorValidacion):
            self.asiento("TRASPASO", "2026-01-01", "x", ("1.1.01", 100, 0), ("3.1.01", 0, 90))

    def test_rechaza_debe_y_haber_en_misma_linea(self):
        with self.assertRaises(ErrorValidacion):
            self.asiento("TRASPASO", "2026-01-01", "x", ("1.1.01", 100, 100), ("3.1.01", 0, 0))

    def test_rechaza_cuenta_de_grupo(self):
        with self.assertRaises(ErrorValidacion):
            self.asiento("TRASPASO", "2026-01-01", "x", ("1.1", 100, 0), ("3.1.01", 0, 100))

    def test_numeracion_por_tipo_y_anio(self):
        a = self.asiento("INGRESO", "2026-01-05", "a", ("1.1.01", 1, 0), ("4.1.01", 0, 1))
        b = self.asiento("INGRESO", "2026-01-06", "b", ("1.1.01", 1, 0), ("4.1.01", 0, 1))
        c = self.asiento("EGRESO", "2026-01-06", "c", ("5.2.12", 1, 0), ("1.1.01", 0, 1))
        d = self.asiento("INGRESO", "2027-01-01", "d", ("1.1.01", 1, 0), ("4.1.01", 0, 1))
        numeros = [nucleo.obtener_asiento(self.conn, self.emp, i)["asiento"]["numero"] for i in (a, b, c, d)]
        self.assertEqual(numeros, [1, 2, 1, 1])


class PruebasInformes(BaseConDatos):
    def test_balance_ocho_columnas_cuadra(self):
        self.cargar_ejemplo()
        b = nucleo.balance_ocho_columnas(self.conn, self.emp, "2026-01-01", "2026-12-31")
        t, s = b["totales"], b["sumas_iguales"]
        self.assertEqual(t["debitos"], t["creditos"])
        self.assertEqual(t["deudor"], t["acreedor"])
        self.assertEqual(s["activo"], s["pasivo"])
        self.assertEqual(s["perdidas"], s["ganancias"])
        self.assertEqual(b["resultado"], 600_000)

    def test_estado_resultados(self):
        self.cargar_ejemplo()
        er = nucleo.estado_resultados(self.conn, self.emp, "2026-01-01", "2026-12-31")
        self.assertEqual(er["total_ingresos"], 1_000_000)
        self.assertEqual(er["total_gastos"], 400_000)
        self.assertEqual(er["resultado"], 600_000)

    def test_balance_general_cuadra(self):
        self.cargar_ejemplo()
        bg = nucleo.balance_general(self.conn, self.emp, "2026-12-31")
        self.assertEqual(bg["total_activo"], 10_000_000 + 1_190_000 - 400_000 + 800_000 - 800_000)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])

    def test_libro_mayor_saldos(self):
        self.cargar_ejemplo()
        m = nucleo.libro_mayor(self.conn, self.emp, self.id_cuenta("1.1.02"), "2026-02-01", "2026-12-31")
        self.assertEqual(m["saldo_inicial"], 10_000_000)
        self.assertEqual(m["saldo_final"], 10_000_000 + 1_190_000 - 400_000 - 952_000)

    def test_cierre_y_anio_siguiente(self):
        self.cargar_ejemplo()
        destino = nucleo.cuenta_resultado_sugerida(self.conn, self.emp)
        self.assertEqual(destino, self.id_cuenta("3.1.04"))
        _, resultado = nucleo.cerrar_ejercicio(self.conn, self.emp, 2026, destino)
        self.assertEqual(resultado, 600_000)
        with self.assertRaises(ErrorValidacion):
            nucleo.cerrar_ejercicio(self.conn, self.emp, 2026, destino)
        # El Estado de Resultados de 2026 no queda en cero por el cierre
        self.assertEqual(nucleo.estado_resultados(self.conn, self.emp, "2026-01-01", "2026-12-31")["resultado"], 600_000)
        # En 2027 el resultado de 2026 ya está en patrimonio
        self.asiento("INGRESO", "2027-01-10", "Venta", ("1.1.01", 50_000, 0), ("4.1.01", 0, 50_000))
        bg = nucleo.balance_general(self.conn, self.emp, "2027-12-31")
        self.assertEqual(bg["resultado_anterior"], 0)
        self.assertEqual(bg["resultado"], 50_000)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])
        b8 = nucleo.balance_ocho_columnas(self.conn, self.emp, "2027-01-01", "2027-12-31")
        self.assertEqual(b8["sumas_iguales"]["activo"], b8["sumas_iguales"]["pasivo"])

    def test_cierre_rechaza_cuenta_que_no_es_patrimonio(self):
        self.cargar_ejemplo()
        with self.assertRaises(ErrorValidacion):
            nucleo.cerrar_ejercicio(self.conn, self.emp, 2026, self.id_cuenta("1.1.01"))

    def test_sin_cierre_arrastra_resultado_anterior(self):
        self.cargar_ejemplo()
        bg = nucleo.balance_general(self.conn, self.emp, "2027-06-30")
        self.assertEqual(bg["resultado_anterior"], 600_000)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])
        b8 = nucleo.balance_ocho_columnas(self.conn, self.emp, "2027-01-01", "2027-12-31")
        self.assertEqual(b8["totales"]["debitos"], b8["totales"]["creditos"])
        self.assertEqual(b8["sumas_iguales"]["activo"], b8["sumas_iguales"]["pasivo"])


class PruebasPatrimonioEnResultados(BaseConDatos):
    """Utilidades acumuladas: en el Balance de 8 columnas van en Ganancias;
    en el Balance General, en el patrimonio."""

    def setUp(self):
        super().setUp()
        cuenta = nucleo.obtener_cuenta(self.conn, self.emp, self.id_cuenta("3.1.02"))
        nucleo.guardar_cuenta(self.conn, self.emp, cuenta["codigo"], cuenta["nombre"], "PATRIMONIO", True, True,
                              cuenta_id=cuenta["id"], resultados_en_ocho=True)
        self.asiento("TRASPASO", "2026-01-02", "Apertura", ("1.1.02", 800_000, 0),
                     ("3.1.01", 0, 300_000), ("3.1.02", 0, 500_000), apertura=True)
        self.asiento("INGRESO", "2026-03-01", "Venta", ("1.1.02", 100_000, 0), ("4.1.01", 0, 100_000))

    def test_ocho_columnas_en_ganancias(self):
        b8 = nucleo.balance_ocho_columnas(self.conn, self.emp, "2026-01-01", "2026-12-31")
        fila = next(f for f in b8["filas"] if f["codigo"] == "3.1.02")
        self.assertEqual((fila["ganancias"], fila["pasivo"]), (500_000, 0))
        self.assertEqual(b8["resultado_ejercicio"], 100_000)
        self.assertEqual(b8["patrimonio_en_resultados"], 500_000)
        self.assertEqual(b8["resultado"], 600_000)
        s = b8["sumas_iguales"]
        self.assertEqual(s["activo"], s["pasivo"])
        self.assertEqual(s["perdidas"], s["ganancias"])

    def test_balance_general_en_patrimonio(self):
        bg = nucleo.balance_general(self.conn, self.emp, "2026-12-31")
        self.assertIn("3.1.02", [f["codigo"] for f in bg["patrimonio"]])
        self.assertEqual(bg["resultado"], 100_000)
        self.assertEqual(bg["total_patrimonio"], 900_000)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])
        er = nucleo.estado_resultados(self.conn, self.emp, "2026-01-01", "2026-12-31")
        self.assertEqual(er["total_ingresos"], 100_000)

    def test_solo_cuentas_de_patrimonio(self):
        with self.assertRaises(ErrorValidacion):
            nucleo.guardar_cuenta(self.conn, self.emp, "1.1.13", "Caja chica", "ACTIVO", True, True,
                                  resultados_en_ocho=True)


class PruebasCierreDeMeses(BaseConDatos):
    def setUp(self):
        super().setUp()
        self.cargar_ejemplo()  # enero a marzo de 2026
        nucleo.cerrar_hasta(self.conn, self.emp, 2026, 2)  # enero y febrero cerrados
        self.venta = self.conn.execute(
            "SELECT id FROM asientos WHERE glosa = 'Venta factura 1'").fetchone()[0]  # 10/02/2026

    def cambios(self):
        return self.conn.execute("SELECT accion, periodos, motivo FROM cambios_periodo_cerrado").fetchall()

    def test_estado_de_los_meses(self):
        meses = nucleo.resumen_periodos(self.conn, self.emp, 2026)
        self.assertEqual([bool(m["cerrado_en"]) for m in meses[:4]], [True, True, False, False])
        self.assertEqual(meses[1]["comprobantes"], 2)
        nucleo.reabrir_mes(self.conn, self.emp, 2026, 1)
        self.assertEqual(nucleo.periodos_cerrados(self.conn, self.emp), {"2026-02"})
        with self.assertRaises(ErrorValidacion):
            nucleo.cerrar_mes(self.conn, self.emp, 2026, 13)

    def test_crear_en_mes_cerrado_pide_confirmacion(self):
        with self.assertRaises(nucleo.PeriodoCerrado) as e:
            self.asiento("TRASPASO", "2026-02-20", "Ajuste", ("1.1.01", 100, 0), ("3.1.01", 0, 100))
        self.assertIn("Febrero 2026 está cerrado", str(e.exception))
        lineas = [{"cuenta_id": self.id_cuenta("1.1.01"), "debe": 100, "haber": 0},
                  {"cuenta_id": self.id_cuenta("3.1.01"), "debe": 0, "haber": 100}]
        nucleo.guardar_asiento(self.conn, self.emp, "TRASPASO", "2026-02-20", "Ajuste", lineas,
                               confirmado=True, motivo="Ajuste solicitado")
        self.assertEqual([tuple(c) for c in self.cambios()], [("Creación", "febrero 2026", "Ajuste solicitado")])

    def test_mover_desde_mes_abierto_a_cerrado_y_al_reves(self):
        compra = self.conn.execute("SELECT id FROM asientos WHERE glosa = 'Compra computador'").fetchone()[0]  # marzo
        datos = nucleo.obtener_asiento(self.conn, self.emp, compra)
        lineas = [dict(l) for l in datos["lineas"]]
        with self.assertRaises(nucleo.PeriodoCerrado):  # de marzo (abierto) a enero (cerrado)
            nucleo.guardar_asiento(self.conn, self.emp, "EGRESO", "2026-01-15", "Compra computador", lineas,
                                   asiento_id=compra)
        venta = nucleo.obtener_asiento(self.conn, self.emp, self.venta)
        with self.assertRaises(nucleo.PeriodoCerrado):  # de febrero (cerrado) a marzo (abierto)
            nucleo.guardar_asiento(self.conn, self.emp, "INGRESO", "2026-03-10", "Venta factura 1",
                                   [dict(l) for l in venta["lineas"]], asiento_id=self.venta)
        # Sin confirmación no cambió nada
        self.assertEqual(nucleo.obtener_asiento(self.conn, self.emp, self.venta)["asiento"]["fecha"], "2026-02-10")
        self.assertEqual(self.cambios(), [])

    def test_mes_abierto_no_pide_nada(self):
        self.asiento("TRASPASO", "2026-03-20", "Ajuste marzo", ("1.1.01", 100, 0), ("3.1.01", 0, 100))
        self.assertEqual(self.cambios(), [])

    def test_eliminar_en_mes_cerrado(self):
        with self.assertRaises(nucleo.PeriodoCerrado):
            nucleo.eliminar_asiento(self.conn, self.emp, self.venta)
        self.assertIsNotNone(nucleo.obtener_asiento(self.conn, self.emp, self.venta))
        nucleo.eliminar_asiento(self.conn, self.emp, self.venta, confirmado=True, motivo="Duplicado")
        self.assertIsNone(nucleo.obtener_asiento(self.conn, self.emp, self.venta))
        cambio = self.conn.execute("SELECT * FROM cambios_periodo_cerrado").fetchone()
        self.assertEqual((cambio["accion"], cambio["motivo"]), ("Eliminación", "Duplicado"))
        self.assertIn("I-1 del 10/02/2026 · Venta factura 1", cambio["comprobante"])

    def test_cierre_de_ejercicio_con_diciembre_cerrado(self):
        nucleo.cerrar_mes(self.conn, self.emp, 2026, 12)
        destino = nucleo.cuenta_resultado_sugerida(self.conn, self.emp)
        with self.assertRaises(nucleo.PeriodoCerrado):
            nucleo.cerrar_ejercicio(self.conn, self.emp, 2026, destino)
        nucleo.cerrar_ejercicio(self.conn, self.emp, 2026, destino, confirmado=True)
        self.assertIsNotNone(nucleo.asiento_cierre(self.conn, self.emp, 2026))

    def test_flujo_web(self):
        cliente = self.app.test_client()
        # Página de períodos y acciones
        texto = cliente.get("/periodos?anio=2026").get_data(as_text=True)
        self.assertIn("Cerrado", texto)
        cliente.post("/periodos/cerrar", data={"anio": 2026, "mes": 3})
        self.assertIn("2026-03", nucleo.periodos_cerrados(self.conn, self.emp))
        cliente.post("/periodos/reabrir", data={"anio": 2026, "mes": 3})
        self.assertNotIn("2026-03", nucleo.periodos_cerrados(self.conn, self.emp))
        self.assertEqual(cliente.post("/periodos/otra", data={"anio": 2026, "mes": 3}).status_code, 404)
        # Editar un comprobante de un mes cerrado: aviso previo, luego confirmación
        texto = cliente.get(f"/asientos/{self.venta}/editar").get_data(as_text=True)
        self.assertIn("pertenece a un mes cerrado (febrero 2026)", texto)
        datos = {"tipo": "INGRESO", "fecha": "2026-02-10", "glosa": "Venta factura 1 corregida",
                 "cuenta_id": [str(self.id_cuenta(c)) for c in ("1.1.02", "4.1.01", "2.1.04")],
                 "glosa_linea": ["", "", ""], "debe": ["1.190.000", "", ""], "haber": ["", "1.000.000", "190.000"]}
        r = cliente.post(f"/asientos/{self.venta}/editar", data=datos)
        self.assertEqual(r.status_code, 200)
        texto = r.get_data(as_text=True)
        self.assertIn("Febrero 2026 está cerrado", texto)
        self.assertIn("Confirmar y guardar", texto)
        self.assertIn("Venta factura 1 corregida", texto)  # conserva lo ingresado
        self.assertEqual(nucleo.obtener_asiento(self.conn, self.emp, self.venta)["asiento"]["glosa"], "Venta factura 1")
        r = cliente.post(f"/asientos/{self.venta}/editar",
                         data=datos | {"confirmar_cerrado": "1", "motivo_cerrado": "Error en glosa"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(nucleo.obtener_asiento(self.conn, self.emp, self.venta)["asiento"]["glosa"],
                         "Venta factura 1 corregida")
        # Eliminar: pide confirmación y luego elimina
        r = cliente.post(f"/asientos/{self.venta}/eliminar")
        self.assertIn("Confirmar eliminación", r.get_data(as_text=True))
        cliente.post(f"/asientos/{self.venta}/eliminar", data={"confirmar_cerrado": "1"})
        self.assertIsNone(nucleo.obtener_asiento(self.conn, self.emp, self.venta))
        # Registro visible en la página
        texto = cliente.get("/periodos?anio=2026").get_data(as_text=True)
        self.assertIn("Error en glosa", texto)
        self.assertIn("Eliminación", texto)
        self.assertIn("mes cerrado", cliente.get("/asientos").get_data(as_text=True))


class PruebasApertura(BaseConDatos):
    """Empresas que reabren sus cuentas cada año con un asiento de apertura."""

    def cargar(self):
        self.asiento("TRASPASO", "2021-09-24", "Aporte de capital", ("1.1.01", 2_000_000, 0), ("3.1.01", 0, 2_000_000))
        self.asiento("TRASPASO", "2021-09-24", "Depósito", ("1.1.01", 0, 2_000_000), ("1.1.02", 2_000_000, 0))
        self.asiento("INGRESO", "2021-11-10", "Venta 2021", ("1.1.02", 300_000, 0), ("4.1.01", 0, 300_000))
        # Apertura 2022: reingresa los saldos, incluido el resultado de 2021 en utilidades retenidas
        self.asiento("TRASPASO", "2022-01-02", "Apertura 2022", ("1.1.02", 2_300_000, 0),
                     ("3.1.01", 0, 2_000_000), ("3.1.02", 0, 300_000), apertura=True)
        self.asiento("INGRESO", "2022-03-01", "Venta 2022", ("1.1.02", 100_000, 0), ("4.1.01", 0, 100_000))

    def test_balance_no_duplica_saldos(self):
        self.cargar()
        bg = nucleo.balance_general(self.conn, self.emp, "2022-12-31")
        self.assertEqual(bg["total_activo"], 2_400_000)
        self.assertEqual(bg["resultado"], 100_000)
        self.assertEqual(bg["resultado_anterior"], 0)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])
        b8 = nucleo.balance_ocho_columnas(self.conn, self.emp, "2022-01-01", "2022-12-31")
        self.assertEqual(b8["sumas_iguales"]["activo"], b8["sumas_iguales"]["pasivo"])
        self.assertEqual(b8["resultado"], 100_000)

    def test_anio_anterior_a_la_apertura(self):
        self.cargar()
        bg = nucleo.balance_general(self.conn, self.emp, "2021-12-31")
        self.assertEqual(bg["total_activo"], 2_300_000)
        self.assertEqual(bg["resultado"], 300_000)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])

    def destinos(self):
        return nucleo.cuentas_destino_sugeridas(self.conn, self.emp)

    def test_propuesta_de_apertura(self):
        self.cargar()
        d = self.destinos()
        self.assertEqual(d["utilidad"], self.id_cuenta("3.1.02"))  # Utilidades retenidas
        self.assertEqual(d["perdida"], self.id_cuenta("3.1.03"))   # Pérdidas acumuladas
        p = nucleo.propuesta_apertura(self.conn, self.emp, 2023, d["utilidad"], d["perdida"])
        lineas = {nucleo.obtener_cuenta(self.conn, self.emp, l["cuenta_id"])["codigo"]: (l["debe"], l["haber"])
                  for l in p["lineas"]}
        # Banco 2.400.000; Capital 2.000.000; Utilidades 300.000 (2021) + 100.000 (utilidad 2022)
        self.assertEqual(lineas, {"1.1.02": (2_400_000, 0), "3.1.01": (0, 2_000_000), "3.1.02": (0, 400_000)})
        self.assertEqual((p["total"], p["resultado"]), (2_400_000, 100_000))
        self.assertEqual(p["glosa"], "APERTURA CUENTAS DE BALANCE AÑO COMERCIAL 2023")
        self.assertEqual(p["existentes"], [])
        # Nada se guardó
        self.assertEqual(nucleo.aperturas_del_anio(self.conn, self.emp, 2023), [])

    def test_propuesta_con_perdida(self):
        self.cargar()
        self.asiento("EGRESO", "2022-06-01", "Gasto grande", ("5.2.12", 500_000, 0), ("1.1.02", 0, 500_000))
        d = self.destinos()
        p = nucleo.propuesta_apertura(self.conn, self.emp, 2023, d["utilidad"], d["perdida"])
        self.assertEqual(p["resultado"], -400_000)  # 100.000 - 500.000
        perdidas = next(l for l in p["lineas"] if l["cuenta_id"] == d["perdida"])
        self.assertEqual((perdidas["debe"], perdidas["haber"]), (400_000, 0))
        self.assertEqual(sum(l["debe"] for l in p["lineas"]), sum(l["haber"] for l in p["lineas"]))

    def test_propuesta_salta_anios_sin_movimiento_y_sin_saldos(self):
        self.cargar()
        d = self.destinos()
        p = nucleo.propuesta_apertura(self.conn, self.emp, 2024, d["utilidad"], d["perdida"])  # 2023 vacío
        self.assertEqual(p["total"], 2_400_000)
        with self.assertRaises(ErrorValidacion):
            nucleo.propuesta_apertura(self.conn, self.emp, 2020, d["utilidad"], d["perdida"])
        with self.assertRaises(ErrorValidacion):  # destino que no es de patrimonio
            nucleo.propuesta_apertura(self.conn, self.emp, 2023, self.id_cuenta("1.1.01"), d["perdida"])

    def test_flujo_web_de_apertura(self):
        self.cargar()
        cliente = self.app.test_client()
        texto = cliente.get("/apertura?anio=2023").get_data(as_text=True)
        self.assertIn("Saldos al 31/12/2022", texto)
        self.assertIn("2.400.000", texto)
        # La propuesta se muestra en el formulario, editable, sin guardar
        r = cliente.get("/apertura/propuesta?anio=2023&fecha=2023-01-02")
        texto = r.get_data(as_text=True)
        self.assertIn("aún no se ha guardado", texto)
        self.assertIn('value="2.400.000"', texto)
        self.assertIn('action="/asientos/nuevo"', texto)
        self.assertEqual(nucleo.aperturas_del_anio(self.conn, self.emp, 2023), [])
        # El usuario modifica un monto y una cuenta antes de guardar
        datos = {"tipo": "TRASPASO", "fecha": "2023-01-02", "glosa": "APERTURA CUENTAS DE BALANCE AÑO COMERCIAL 2023",
                 "apertura": "on",
                 "cuenta_id": [str(self.id_cuenta(c)) for c in ("1.1.02", "1.1.01", "3.1.01", "3.1.02")],
                 "glosa_linea": [""] * 4, "debe": ["2.300.000", "100.000", "", ""],
                 "haber": ["", "", "2.000.000", "400.000"]}
        self.assertEqual(cliente.post("/asientos/nuevo", data=datos).status_code, 302)
        abiertas = nucleo.aperturas_del_anio(self.conn, self.emp, 2023)
        self.assertEqual(len(abiertas), 1)
        bg = nucleo.balance_general(self.conn, self.emp, "2023-12-31")
        self.assertEqual(bg["total_activo"], 2_400_000)
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])
        # Si ya hay una apertura, se advierte
        self.assertIn("Ya existe un asiento de apertura en 2023",
                      cliente.get("/apertura/propuesta?anio=2023").get_data(as_text=True))
        # Un año sin saldos vuelve a la página con el error
        r = cliente.get("/apertura/propuesta?anio=2020", follow_redirects=True)
        self.assertIn("No hay saldos", r.get_data(as_text=True))

    def test_mayor_reinicia_saldo_en_apertura(self):
        self.cargar()
        m = nucleo.libro_mayor(self.conn, self.emp, self.id_cuenta("1.1.02"), "2021-01-01", "2022-12-31")
        self.assertEqual(m["saldo_final"], 2_400_000)
        self.assertTrue(any(x.get("reinicio") for x in m["movimientos"]))
        m2022 = nucleo.libro_mayor(self.conn, self.emp, self.id_cuenta("1.1.02"), "2022-01-01", "2022-12-31")
        self.assertEqual(m2022["saldo_final"], 2_400_000)


class PruebasCuentas(BaseConDatos):
    def test_crear_subcuenta_hereda_tipo(self):
        cid = nucleo.guardar_cuenta(self.conn, self.emp, "1.1.12", "Fondo fijo", "", True, True)
        self.assertEqual(nucleo.obtener_cuenta(self.conn, self.emp, cid)["tipo"], "ACTIVO")

    def test_tipo_explicito(self):
        cid = nucleo.guardar_cuenta(self.conn, self.emp, "4.2.03", "Pérdida en venta de activos", "GASTO", True, True)
        self.assertEqual(nucleo.obtener_cuenta(self.conn, self.emp, cid)["tipo"], "GASTO")
        self.asiento("TRASPASO", "2026-05-01", "x", ("4.2.03", 1000, 0), ("1.1.01", 0, 1000))
        er = nucleo.estado_resultados(self.conn, self.emp, "2026-01-01", "2026-12-31")
        self.assertEqual(er["total_gastos"], 1000)
        self.assertIn("INGRESOS NO OPERACIONALES", [f["nombre"] for f in er["gastos"]])

    def test_no_crea_sin_padre_ni_bajo_imputable(self):
        with self.assertRaises(ErrorValidacion):
            nucleo.guardar_cuenta(self.conn, self.emp, "7.9.01", "Sin padre", "ACTIVO", True, True)
        with self.assertRaises(ErrorValidacion):
            nucleo.guardar_cuenta(self.conn, self.emp, "1.1.01.01", "Bajo imputable", "", True, True)

    def test_no_elimina_con_movimientos(self):
        self.cargar_ejemplo()
        with self.assertRaises(ErrorValidacion):
            nucleo.eliminar_cuenta(self.conn, self.emp, self.id_cuenta("1.1.02"))


class PruebasEmpresas(BaseConDatos):
    def test_datos_aislados_entre_empresas(self):
        otra = self.nueva_empresa()
        self.cargar_ejemplo()
        self.asiento("INGRESO", "2026-04-01", "Venta otra", ("1.1.01", 5_000, 0), ("4.1.01", 0, 5_000), empresa=otra)
        self.assertEqual(nucleo.estado_resultados(self.conn, self.emp, "2026-01-01", "2026-12-31")["total_ingresos"], 1_000_000)
        self.assertEqual(nucleo.estado_resultados(self.conn, otra, "2026-01-01", "2026-12-31")["total_ingresos"], 5_000)
        # No se puede usar una cuenta de otra empresa
        with self.assertRaises(ErrorValidacion):
            nucleo.guardar_asiento(self.conn, otra, "TRASPASO", "2026-01-01", "x",
                                   [{"cuenta_id": self.id_cuenta("1.1.01"), "debe": 1, "haber": 0},
                                    {"cuenta_id": self.id_cuenta("3.1.01", otra), "debe": 0, "haber": 1}])
        # Ni ver un asiento de otra empresa
        primero = self.conn.execute("SELECT MIN(id) FROM asientos WHERE empresa_id = ?", (self.emp,)).fetchone()[0]
        self.assertIsNone(nucleo.obtener_asiento(self.conn, otra, primero))

    def test_rut_repetido(self):
        self.nueva_empresa()
        with self.assertRaises(ErrorValidacion):
            self.nueva_empresa("Copia SpA", "11111111-1")

    def test_eliminar_empresa(self):
        otra = self.nueva_empresa()
        self.cargar_ejemplo(empresa=otra)
        nucleo.eliminar_empresa(self.conn, otra)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0], 0)
        with self.assertRaises(ErrorValidacion):
            nucleo.eliminar_empresa(self.conn, self.emp)  # la única


class PruebasMigracion(unittest.TestCase):
    def test_base_de_una_empresa_se_migra(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "vieja.db"
            vieja = sqlite3.connect(ruta)
            vieja.executescript("""
                CREATE TABLE empresa (id INTEGER PRIMARY KEY CHECK (id = 1), razon_social TEXT NOT NULL DEFAULT '',
                  rut TEXT NOT NULL DEFAULT '', giro TEXT NOT NULL DEFAULT '', direccion TEXT NOT NULL DEFAULT '',
                  comuna TEXT NOT NULL DEFAULT '', representante TEXT NOT NULL DEFAULT '', ejercicio INTEGER NOT NULL);
                CREATE TABLE cuentas (id INTEGER PRIMARY KEY, codigo TEXT NOT NULL UNIQUE, nombre TEXT NOT NULL,
                  tipo TEXT NOT NULL, imputable INTEGER NOT NULL DEFAULT 1, activa INTEGER NOT NULL DEFAULT 1);
                CREATE TABLE asientos (id INTEGER PRIMARY KEY, tipo TEXT NOT NULL, numero INTEGER NOT NULL,
                  fecha TEXT NOT NULL, glosa TEXT NOT NULL, cierre INTEGER NOT NULL DEFAULT 0,
                  creado_en TEXT NOT NULL DEFAULT (datetime('now', 'localtime')));
                CREATE INDEX idx_asientos_fecha ON asientos (fecha);
                CREATE TABLE lineas (id INTEGER PRIMARY KEY, asiento_id INTEGER NOT NULL REFERENCES asientos (id) ON DELETE CASCADE,
                  orden INTEGER NOT NULL, cuenta_id INTEGER NOT NULL REFERENCES cuentas (id), glosa TEXT NOT NULL DEFAULT '',
                  debe INTEGER NOT NULL DEFAULT 0, haber INTEGER NOT NULL DEFAULT 0);
                INSERT INTO empresa VALUES (1, 'Antigua SpA', '12.345.678-5', '', '', '', '', 2026);
                INSERT INTO cuentas VALUES (1, '1', 'ACTIVO', 'ACTIVO', 0, 1), (2, '1.1', 'CAJA', 'ACTIVO', 1, 1),
                                           (3, '3', 'PATRIMONIO', 'PATRIMONIO', 0, 1), (4, '3.1', 'CAPITAL', 'PATRIMONIO', 1, 1);
                INSERT INTO asientos (id, tipo, numero, fecha, glosa) VALUES (1, 'TRASPASO', 1, '2026-01-01', 'Apertura');
                INSERT INTO lineas VALUES (1, 1, 1, 2, '', 500, 0), (2, 1, 2, 4, '', 0, 500);
            """)
            vieja.commit()
            vieja.close()
            conn = db.conectar(ruta)
            db.inicializar(conn)
            empresa = conn.execute("SELECT * FROM empresas").fetchone()
            self.assertEqual(empresa["razon_social"], "Antigua SpA")
            bg = nucleo.balance_general(conn, empresa["id"], "2026-12-31")
            self.assertEqual(bg["total_activo"], 500)
            self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])
            conn.close()


def crear_origen_creandogestion(ruta):
    """Base mínima con el formato de CreandoGestion para probar el importador."""
    c = sqlite3.connect(ruta)
    c.executescript("""
        CREATE TABLE companies (id TEXT PRIMARY KEY, name TEXT, rut TEXT, giro TEXT, direccion TEXT,
                                ciudad TEXT, regime TEXT, accounting TEXT);
        CREATE TABLE accounts (id TEXT PRIMARY KEY, code TEXT, name TEXT, type TEXT, active INTEGER, company_id TEXT);
        CREATE TABLE entries (id TEXT PRIMARY KEY, number INTEGER, date TEXT, description TEXT, company_id TEXT);
        CREATE TABLE entry_details (id TEXT PRIMARY KEY, entry_id TEXT, account_code TEXT, debit REAL, credit REAL);
        INSERT INTO companies VALUES ('c1', 'PRUEBA IMPORTADA SPA', '12.345.678-5', 'Servicios', 'Calle 1', 'Santiago', '14D3', 'completa');
        INSERT INTO accounts VALUES
          ('a1', '1.1.10.1', 'CAJA', 'Activo', 1, 'c1'), ('a2', '1.1.20.1', 'BANCO', 'Activo', 1, 'c1'),
          ('a3', '3.1.10.1', 'CAPITAL PAGADO', 'Patrimonio', 1, 'c1'),
          ('a4', '3.2.10.1', 'UTILIDADES ACUMULADAS', 'Ganancia', 1, 'c1'),
          ('a5', '4.1.10.1', 'VENTAS', 'Ganancia', 1, 'c1'), ('a6', '4.3.10.1', 'GASTOS GENERALES', 'Pérdida', 1, 'c1'),
          ('a7', '4.5.10.1', 'GASTOS FINANCIEROS', 'Pérdida', 1, 'c1'), ('a8', '4.5.20.1', 'OTROS INGRESOS', 'Ganancia', 1, 'c1');
        INSERT INTO entries VALUES
          ('e1', 1, '2021-09-24', 'APORTE DE CAPITAL', 'c1'),
          ('e2', 2, '2021-10-01', 'VENTA', 'c1'),
          ('e3', 3, '2022-01-02', 'APERTURA CUENTAS DE BALANCE 2022', 'c1'),
          ('e4', 4, '2022-02-01', 'GASTOS Y OTROS', 'c1');
        INSERT INTO entry_details VALUES
          ('d1', 'e1', '1.1.20.1', 1000000, 0), ('d2', 'e1', '3.1.10.1', 0, 1000000),
          ('d3', 'e2', '1.1.20.1', 50000, 0), ('d4', 'e2', '4.1.10.1', 0, 50000),
          ('d5', 'e3', '1.1.20.1', 1050000, 0), ('d6', 'e3', '3.1.10.1', 0, 1000000), ('d7', 'e3', '3.2.10.1', 0, 50000),
          ('d8', 'e4', '4.3.10.1', 20000, 0), ('d9', 'e4', '4.5.10.1', 5000, 0), ('d10', 'e4', '4.5.20.1', 0, 3000),
          ('d11', 'e4', '1.1.20.1', 0, 22000);
    """)
    c.commit()
    c.close()


class PruebasImportador(BaseConDatos):
    def setUp(self):
        super().setUp()
        self.origen = Path(self.tmp.name) / "origen.db"
        crear_origen_creandogestion(self.origen)

    def test_importa_y_cuadra(self):
        resultados, eliminadas = importar_creandogestion.importar(self.conn, self.origen, ["c1"])
        r = resultados[0]
        self.assertEqual(eliminadas, 1)  # la empresa inicial vacía
        self.assertEqual((r["asientos"], r["aperturas"], r["ejercicio"]), (4, 1, 2022))
        # 3.2.10.1 venía como «Ganancia»: patrimonio, pero en Resultados del Balance de 8 columnas
        self.assertEqual(len(r["correcciones"]), 1)
        e = r["empresa_id"]
        utilidades = self.conn.execute(
            "SELECT tipo, resultados_en_ocho FROM cuentas WHERE empresa_id = ? AND codigo = '3.2.10.1'", (e,)
        ).fetchone()
        self.assertEqual(tuple(utilidades), ("PATRIMONIO", 1))
        b8 = nucleo.balance_ocho_columnas(self.conn, e, "2022-01-01", "2022-12-31")
        self.assertEqual(b8["patrimonio_en_resultados"], 50_000)
        bg = nucleo.balance_general(self.conn, e, "2022-12-31")
        self.assertEqual(bg["total_activo"], 1_050_000 - 22_000)  # sin duplicar la apertura
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])
        er = nucleo.estado_resultados(self.conn, e, "2022-01-01", "2022-12-31")
        self.assertEqual((er["total_ingresos"], er["total_gastos"]), (3_000, 25_000))
        nombres = [f["nombre"] for f in er["gastos"] if f["grupo"]]
        self.assertIn("GASTOS DE ADMINISTRACIÓN Y VENTAS", nombres)
        self.assertIn("OTROS RESULTADOS NO OPERACIONALES", nombres)

    def test_no_importa_dos_veces(self):
        importar_creandogestion.importar(self.conn, self.origen, ["c1"])
        antes = self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0]
        with self.assertRaises(ErrorValidacion):
            importar_creandogestion.importar(self.conn, self.origen, ["c1"])
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0], antes)

    def test_falla_sin_dejar_datos_a_medias(self):
        c = sqlite3.connect(self.origen)
        c.execute("UPDATE entry_details SET credit = 999 WHERE id = 'd11'")  # asiento 4 descuadrado
        c.commit()
        c.close()
        with self.assertRaises(ErrorValidacion):
            importar_creandogestion.importar(self.conn, self.origen, ["c1"])
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM empresas").fetchone()[0], 1)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0], 0)


class PruebasWeb(BaseConDatos):
    def test_todas_las_paginas_cargan(self):
        self.cargar_ejemplo()
        self.nueva_empresa()
        cliente = self.app.test_client()
        asiento_id = self.conn.execute("SELECT MIN(id) FROM asientos").fetchone()[0]
        urls = ["/", "/empresa", "/empresas", "/empresas/nueva", "/cuentas", "/cuentas/nueva",
                f"/cuentas/{self.id_cuenta('1.1.01')}/editar", "/cuentas/nueva?codigo=4.2.",
                "/asientos", "/asientos?tipo=INGRESO&q=Venta", "/asientos/nuevo", f"/asientos/{asiento_id}",
                f"/asientos/{asiento_id}/editar", "/libro-diario?desde=2026-01-01&hasta=2026-12-31",
                f"/libro-mayor?cuenta={self.id_cuenta('1.1.02')}&desde=2026-01-01&hasta=2026-12-31",
                "/libro-mayor", "/balance-8-columnas?desde=2026-01-01&hasta=2026-12-31",
                "/estado-resultados?desde=2026-01-01&hasta=2026-12-31", "/balance-general?hasta=2026-12-31",
                "/cierre?anio=2026", "/libro-diario?formato=csv", "/balance-8-columnas?formato=csv",
                f"/libro-mayor?cuenta={self.id_cuenta('1.1.02')}&formato=csv", "/respaldo", "/respaldos"]
        for url in urls:
            with self.subTest(url=url):
                self.assertEqual(cliente.get(url).status_code, 200)
        self.assertEqual(cliente.get("/asientos/9999").status_code, 404)

    def test_libro_mayor_todas_las_cuentas_en_pantalla(self):
        self.cargar_ejemplo()
        cliente = self.app.test_client()
        r = cliente.get("/libro-mayor?cuenta=todas&desde=2026-01-01&hasta=2026-12-31")
        self.assertEqual(r.status_code, 200)
        texto = r.get_data(as_text=True)
        self.assertEqual(texto.count('class="seccion-mayor"'), 7)  # una sección por cuenta usada
        self.assertIn("7 cuentas con movimiento", texto)
        self.assertIn("(CONCORDADO)", texto)
        self.assertIn("Saldo deudor", texto)    # Banco
        self.assertIn("Saldo acreedor", texto)  # Capital, Ventas…
        self.assertIn('value="todas" selected', texto)
        # Los botones de exportación de esta vista exportan el mayor completo
        r = cliente.get("/libro-mayor?cuenta=todas&desde=2026-01-01&hasta=2026-12-31&formato=csv")
        self.assertIn("(CONCORDADO)", r.data.decode("utf-8-sig"))
        # Sin cuenta elegida, ofrece ver todas
        self.assertIn("Ver todas las cuentas", cliente.get("/libro-mayor").get_data(as_text=True))

    def test_cambiar_de_empresa(self):
        self.cargar_ejemplo()
        otra = self.nueva_empresa("Segunda SpA")
        cliente = self.app.test_client()
        asiento_id = self.conn.execute("SELECT MIN(id) FROM asientos").fetchone()[0]
        cliente.post("/empresas/activar", data={"empresa_id": otra})
        texto = cliente.get("/").get_data(as_text=True)
        self.assertIn("0 en el ejercicio", texto)
        # El comprobante de la primera empresa no es visible desde la segunda
        self.assertEqual(cliente.get(f"/asientos/{asiento_id}").status_code, 404)
        cliente.post("/empresas/activar", data={"empresa_id": self.emp})
        self.assertEqual(cliente.get(f"/asientos/{asiento_id}").status_code, 200)

    def test_crear_comprobante_por_formulario(self):
        cliente = self.app.test_client()
        datos = {"tipo": "TRASPASO", "fecha": "2026-01-01", "glosa": "Apertura", "apertura": "on",
                 "cuenta_id": [str(self.id_cuenta("1.1.01")), str(self.id_cuenta("3.1.01")), ""],
                 "glosa_linea": ["", "", ""], "debe": ["1.500.000", "", ""], "haber": ["", "1.500.000", ""]}
        r = cliente.post("/asientos/nuevo", data=datos)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.conn.execute("SELECT SUM(debe) FROM lineas").fetchone()[0], 1_500_000)
        self.assertEqual(self.conn.execute("SELECT apertura FROM asientos").fetchone()[0], 1)
        # Descuadrado: vuelve al formulario con error
        datos["haber"] = ["", "1.000.000", ""]
        r = cliente.post("/asientos/nuevo", data=datos)
        self.assertEqual(r.status_code, 200)
        self.assertIn("no cuadra", r.get_data(as_text=True))

    def test_empresa_valida_rut(self):
        cliente = self.app.test_client()
        r = cliente.post("/empresa", data={"razon_social": "Mi Pyme SpA", "rut": "12345678-9", "ejercicio": "2026"})
        self.assertIn("RUT no es válido", r.get_data(as_text=True))
        r = cliente.post("/empresa", data={"razon_social": "Mi Pyme SpA", "rut": "123456785", "ejercicio": "2026"})
        self.assertEqual(r.status_code, 302)
        rut = self.conn.execute("SELECT rut FROM empresas WHERE id = ?", (self.emp,)).fetchone()[0]
        self.assertEqual(rut, "12.345.678-5")

    def test_eliminar_empresa_pide_confirmacion(self):
        otra = self.nueva_empresa("Borrable SpA")
        cliente = self.app.test_client()
        cliente.post("/empresas/activar", data={"empresa_id": otra})
        cliente.post("/empresa/eliminar", data={"confirmacion": "otra cosa"})
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM empresas").fetchone()[0], 2)
        cliente.post("/empresa/eliminar", data={"confirmacion": "borrable spa"})
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM empresas").fetchone()[0], 1)


class PruebasExportacion(BaseConDatos):
    def test_todos_los_informes_en_todos_los_formatos(self):
        from openpyxl import load_workbook

        self.cargar_ejemplo()
        cuenta = nucleo.obtener_cuenta(self.conn, self.emp, self.id_cuenta("3.1.02"))
        nucleo.guardar_cuenta(self.conn, self.emp, "3.1.02", cuenta["nombre"], "PATRIMONIO", True, True,
                              cuenta_id=cuenta["id"], resultados_en_ocho=True)
        self.asiento("TRASPASO", "2026-04-01", "Traspaso a utilidades & <prueba>", ("1.1.01", 1000, 0), ("3.1.02", 0, 1000))
        cliente = self.app.test_client()
        cliente.post("/empresa", data={"razon_social": "Mi Pyme SpA", "rut": "123456785", "ejercicio": "2026"})
        asiento_id = self.conn.execute("SELECT MIN(id) FROM asientos").fetchone()[0]
        rango = "desde=2026-01-01&hasta=2026-12-31"
        paginas = {
            "libro_diario": f"/libro-diario?{rango}",
            "mayor": f"/libro-mayor?cuenta={self.id_cuenta('1.1.02')}&{rango}",
            "balance_8_columnas": f"/balance-8-columnas?{rango}",
            "estado_resultados": f"/estado-resultados?{rango}",
            "estado_situacion_financiera": "/balance-general?hasta=2026-12-31",
            "libro_mayor": f"/libro-mayor?todas=1&{rango}",
            "comprobante": f"/asientos/{asiento_id}",
            "plan_de_cuentas": "/cuentas",
        }
        for nombre, url in paginas.items():
            for formato, firma in (("pdf", b"%PDF"), ("xlsx", b"PK"), ("csv", "﻿".encode())):
                with self.subTest(informe=nombre, formato=formato):
                    r = cliente.get(f"{url}&formato={formato}" if "?" in url else f"{url}?formato={formato}")
                    self.assertEqual(r.status_code, 200)
                    self.assertTrue(r.data.startswith(firma))
                    disposicion = r.headers["Content-Disposition"]
                    self.assertIn(f'mi_pyme_spa_{nombre}', disposicion)
                    self.assertTrue(disposicion.endswith(f'.{formato}"'))
                    if formato == "xlsx":
                        hoja = load_workbook(io.BytesIO(r.data)).active
                        self.assertEqual(hoja["A1"].value, "RAZON SOCIAL: MI PYME SPA")

    def test_excel_con_numeros_reales(self):
        from openpyxl import load_workbook

        self.cargar_ejemplo()
        cliente = self.app.test_client()
        r = cliente.get("/balance-8-columnas?desde=2026-01-01&hasta=2026-12-31&formato=xlsx")
        hoja = load_workbook(io.BytesIO(r.data)).active
        filas = {fila[1]: fila for fila in hoja.iter_rows(values_only=True) if fila and len(fila) > 2}
        self.assertEqual(filas["BANCO"][2], 11_190_000)  # débitos como número, no texto
        self.assertEqual(filas["TOTALES GENERALES:"][6], filas["TOTALES GENERALES:"][7])  # activo = pasivo

    def pdf(self, cliente, url):
        r = cliente.get(url)
        self.assertEqual(r.status_code, 200)
        return r

    def test_folios_correlativos(self):
        self.cargar_ejemplo()
        cliente = self.app.test_client()
        rango = "desde=2026-01-01&hasta=2026-12-31"
        r1 = self.pdf(cliente, f"/balance-8-columnas?{rango}&formato=pdf")
        self.assertEqual(r1.headers["X-Folios"], "1-1")
        r2 = self.pdf(cliente, f"/libro-diario?{rango}&formato=pdf")
        self.assertEqual(r2.headers["X-Folios"], "2-2")
        # Los comprobantes, el Excel y el CSV no usan folios
        asiento_id = self.conn.execute("SELECT MIN(id) FROM asientos").fetchone()[0]
        self.assertNotIn("X-Folios", self.pdf(cliente, f"/asientos/{asiento_id}?formato=pdf").headers)
        self.pdf(cliente, f"/libro-diario?{rango}&formato=xlsx")
        self.assertEqual(self.conn.execute("SELECT ultimo_folio FROM empresas WHERE id = ?", (self.emp,)).fetchone()[0], 2)
        # El folio se puede ajustar desde los datos de la empresa
        cliente.post("/empresa", data={"razon_social": "Mi Pyme SpA", "ejercicio": "2026", "ultimo_folio": "39"})
        self.assertEqual(self.pdf(cliente, f"/estado-resultados?{rango}&formato=pdf").headers["X-Folios"], "40-40")
        # Cada empresa lleva su propia numeración
        otra = self.nueva_empresa()
        cliente.post("/empresas/activar", data={"empresa_id": otra})
        self.assertEqual(self.pdf(cliente, f"/balance-general?hasta=2026-12-31&formato=pdf").headers["X-Folios"], "1-1")

    def test_folio_inicial_indicado(self):
        self.cargar_ejemplo()
        cliente = self.app.test_client()
        rango = "desde=2026-01-01&hasta=2026-12-31"
        ultimo = lambda: self.conn.execute("SELECT ultimo_folio FROM empresas WHERE id = ?", (self.emp,)).fetchone()[0]
        r = self.pdf(cliente, f"/libro-diario?{rango}&formato=pdf&folio=100")
        self.assertEqual(r.headers["X-Folios"], "100-100")
        self.assertEqual(ultimo(), 100)
        # Reimprimir con un folio anterior no hace retroceder la numeración
        self.assertEqual(self.pdf(cliente, f"/libro-diario?{rango}&formato=pdf&folio=40").headers["X-Folios"], "40-40")
        self.assertEqual(ultimo(), 100)
        # Un folio inválido usa el siguiente al último
        self.assertEqual(self.pdf(cliente, f"/libro-diario?{rango}&formato=pdf&folio=0").headers["X-Folios"], "101-101")
        # La página ofrece el diálogo con el folio sugerido y marca los botones de PDF oficiales
        html = cliente.get(f"/libro-diario?{rango}").get_data(as_text=True)
        self.assertIn('data-proximo="102"', html)
        self.assertIn("data-folio", html)
        asiento_id = self.conn.execute("SELECT MIN(id) FROM asientos").fetchone()[0]
        self.assertNotIn("data-folio>", cliente.get(f"/asientos/{asiento_id}").get_data(as_text=True))

    def test_guardar_empresa_sin_campo_folio_no_lo_borra(self):
        self.conn.execute("UPDATE empresas SET ultimo_folio = 25")
        self.conn.commit()
        nucleo.guardar_empresa(self.conn, self.emp, {"razon_social": "X SpA", "ejercicio": "2026"})
        self.assertEqual(self.conn.execute("SELECT ultimo_folio FROM empresas").fetchone()[0], 25)

    def test_contador_y_firma(self):
        from PIL import Image

        cliente = self.app.test_client()
        r = cliente.post("/contador", data={"nombre": "Ana Pérez", "rut": "12345678-9"}, follow_redirects=True)
        self.assertIn("RUT del contador no es válido", r.get_data(as_text=True))
        imagen = io.BytesIO()
        Image.new("RGB", (1200, 400), "white").save(imagen, format="JPEG")
        cliente.post("/contador", data={"nombre": "Ana Pérez", "rut": "123456785",
                                        "firma": (io.BytesIO(imagen.getvalue()), "firma.jpg")},
                     content_type="multipart/form-data")
        contador = nucleo.obtener_contador(self.conn)
        self.assertEqual((contador["nombre"], contador["rut"]), ("Ana Pérez", "12.345.678-5"))
        self.assertEqual(Image.open(io.BytesIO(contador["firma"])).size, (800, 267))  # PNG reducido
        self.assertEqual(cliente.get("/contador/firma.png").status_code, 200)
        # La firma aparece en el PDF y en el Excel
        self.cargar_ejemplo()
        r = cliente.get("/estado-resultados?desde=2026-01-01&hasta=2026-12-31&formato=pdf")
        self.assertIn(b"/Image", r.data)
        r = cliente.get("/estado-resultados?desde=2026-01-01&hasta=2026-12-31&formato=xlsx")
        self.assertIn(b"xl/media/image1.png", r.data)
        r = cliente.post("/contador", data={"nombre": "x", "firma": (io.BytesIO(b"no es imagen"), "f.png")},
                         content_type="multipart/form-data")
        self.assertIn("no es una imagen válida", r.get_data(as_text=True))
        cliente.post("/contador/firma/eliminar")
        self.assertIsNone(nucleo.obtener_contador(self.conn)["firma"])

    def test_libro_mayor_completo_concordado(self):
        self.cargar_ejemplo()
        completo = nucleo.libro_mayor_completo(self.conn, self.emp, "2026-01-01", "2026-12-31")
        diario = nucleo.libro_diario(self.conn, self.emp, "2026-01-01", "2026-12-31")
        self.assertEqual((completo["total_debe"], completo["total_haber"]),
                         (diario["total_debe"], diario["total_haber"]))
        self.assertEqual(len(completo["cuentas"]), 7)  # cuentas usadas en el ejemplo
        r = self.app.test_client().get("/libro-mayor?todas=1&desde=2026-01-01&hasta=2026-12-31&formato=csv")
        self.assertIn("(CONCORDADO)", r.data.decode("utf-8-sig"))

    def test_csv_del_diario_es_plano(self):
        self.cargar_ejemplo()
        r = self.app.test_client().get("/libro-diario?desde=2026-01-01&hasta=2026-12-31&formato=csv")
        lineas = r.data.decode("utf-8-sig").strip().splitlines()
        self.assertEqual(lineas[0], "Fecha;Comprobante;Glosa;Código;Cuenta;Detalle;Debe;Haber")
        self.assertEqual(len(lineas), 1 + 10)  # una fila por línea de comprobante (2 + 3 + 2 + 3)


CSV_COMPRAS = (
    "Nro;Tipo Doc;Tipo Compra;RUT Proveedor;Razon Social;Folio;Fecha Docto;Fecha Recepcion;Fecha Acuse;"
    "Monto Exento;Monto Neto;Monto IVA Recuperable;Monto Iva No Recuperable;Codigo IVA No Rec.;Monto Total;"
    "Monto Neto Activo Fijo;IVA Activo Fijo;IVA uso Comun;Impto. Sin Derecho a Credito;IVA No Retenido;"
    "Tabacos Puros;Tabacos Cigarrillos;Tabacos Elaborados;NCE o NDE sobre Fact. de Compra;Codigo Otro Impuesto;"
    "Valor Otro Impuesto;Tasa Otro Impuesto\n"
    "1;33;Del Giro;76999888-7;Proveedor Uno;1001;10/05/2026;11/05/2026;;0;1000000;190000;0;;1190000;0;0;0;0;0;;;;0;;;\n"
    "2;33;Del Giro;76555444-3;Ferretería Dos;2002;12/05/2026;12/05/2026;;0;170000;32300;0;;202300;0;0;0;0;0;;;;0;;;\n"
    "3;33;Del Giro;77111222-3;Máquinas SpA;3003;15/05/2026;15/05/2026;;0;500000;95000;0;;595000;500000;95000;0;0;0;;;;0;;;\n"
    "4;61;Del Giro;76999888-7;Proveedor Uno;55;20/05/2026;20/05/2026;;0;100000;19000;0;;119000;0;0;0;0;0;;;;0;;;\n"
    "5;33;Del Giro;76693183-9;Combustibles SA;4004;22/05/2026;22/05/2026;;0;50000;9500;0;;65000;0;0;0;0;0;;;;0;28;5500;\n"
    "6;34;Del Giro;78123123-K;Asesorías Exentas;77;25/05/2026;25/05/2026;;40000;0;0;0;;40000;0;0;0;0;0;;;;0;;;\n"
)
CSV_VENTAS = (
    "Nro;Tipo Doc;Tipo Venta;Rut cliente;Razon Social;Folio;Fecha Docto;Fecha Recepcion;Fecha Acuse Recibo;"
    "Fecha Reclamo;Monto Exento;Monto Neto;Monto IVA;Monto total;IVA Retenido Total;IVA Retenido Parcial\n"
    "1;33;Del Giro;88888888-8;Cliente Uno SpA;101;05/05/2026;;;;0;1000000;190000;1190000;0;0\n"
    "2;33;Del Giro;99999999-9;Cliente Dos Ltda;102;18/05/2026;;;;0;500000;95000;595000;0;0\n"
    "3;61;Del Giro;88888888-8;Cliente Uno SpA;5;28/05/2026;;;;0;100000;19000;119000;0;0\n"
    "4;34;Del Giro;99999999-9;Cliente Dos Ltda;103;30/05/2026;;;;200000;0;0;200000;0;0\n"
)


class PruebasRCVyF29(BaseConDatos):
    def importar(self, texto=CSV_COMPRAS, periodo="2026-05", codificacion="utf-8"):
        from contabilidad import tributario
        return tributario.importar_rcv(self.conn, self.emp, periodo, texto.encode(codificacion))

    def configurar_otros_impuestos(self):
        from contabilidad import tributario
        tributario.guardar_cuentas_tributarias(self.conn, self.emp, {"compras_otros_impuestos": self.id_cuenta("5.2.12")},
                                               tasa_ppm="0,25")

    def saldos_asiento(self, asiento_id):
        datos = nucleo.obtener_asiento(self.conn, self.emp, asiento_id)
        return {l["codigo"]: l["debe"] - l["haber"] for l in datos["lineas"]}

    def test_lectura_csv(self):
        from contabilidad import tributario
        registro, docs = tributario.leer_csv_rcv(CSV_COMPRAS.encode("cp1252"))  # también en Latin-1
        self.assertEqual((registro, len(docs)), ("COMPRA", 6))
        self.assertEqual(docs[0]["rut"], "76.999.888-7")
        self.assertEqual(docs[0]["fecha"], "2026-05-10")
        self.assertEqual(docs[1]["razon_social"], "Ferretería Dos")
        self.assertEqual(docs[2]["activo_fijo"], 1)
        self.assertEqual(docs[4]["otros_impuestos"], 5500)
        self.assertEqual(docs[5]["exento"], 40000)
        registro, docs = tributario.leer_csv_rcv(CSV_VENTAS.encode())
        self.assertEqual((registro, len(docs)), ("VENTA", 4))
        with self.assertRaises(ErrorValidacion):
            tributario.leer_csv_rcv(b"columna;otra\n1;2\n")

    def test_importar_sin_duplicar(self):
        self.assertEqual(self.importar(), ("COMPRA", 6, 0))
        self.assertEqual(self.importar(), ("COMPRA", 0, 6))

    def test_centralizacion_de_compras(self):
        from contabilidad import tributario
        self.importar()
        with self.assertRaises(ErrorValidacion) as e:  # falta la cuenta de otros impuestos
            tributario.propuesta_centralizacion(self.conn, self.emp, "2026-05", "COMPRA")
        self.assertIn("otros impuestos", str(e.exception))
        self.configurar_otros_impuestos()
        p = tributario.propuesta_centralizacion(self.conn, self.emp, "2026-05", "COMPRA")
        lineas = {nucleo.obtener_cuenta(self.conn, self.emp, l["cuenta_id"])["codigo"]: l["debe"] - l["haber"]
                  for l in p["lineas"]}
        self.assertEqual(lineas, {"1.1.08": 1_160_000,   # Mercaderías: netos + exento - NC
                                  "1.2.03": 500_000,     # Maquinarias (activo fijo)
                                  "1.1.09": 307_800,     # IVA crédito fiscal
                                  "5.2.12": 5_500,       # impuesto específico
                                  "2.1.01": -1_973_300})  # Proveedores
        self.assertEqual(p["glosa"], "CENTRALIZACIÓN COMPRAS MAYO 2026")
        self.assertEqual(p["fecha"], "2026-05-31")

    def test_cuenta_por_documento_y_recordada_por_rut(self):
        from contabilidad import tributario
        self.importar()
        self.configurar_otros_impuestos()
        doc = self.conn.execute("SELECT id FROM rcv_documentos WHERE folio = '2002'").fetchone()[0]
        tributario.asignar_cuentas(self.conn, self.emp, {doc: self.id_cuenta("5.2.08")})  # Mantención
        p = tributario.propuesta_centralizacion(self.conn, self.emp, "2026-05", "COMPRA")
        montos = {l["cuenta_id"]: l["debe"] for l in p["lineas"]}
        self.assertEqual(montos[self.id_cuenta("5.2.08")], 170_000)
        # El mes siguiente, la misma ferretería ya viene con esa cuenta
        self.importar(CSV_COMPRAS.replace("2002;12/05/2026", "2010;12/06/2026"), periodo="2026-06")
        cuenta = self.conn.execute("SELECT cuenta_id FROM rcv_documentos WHERE folio = '2010'").fetchone()[0]
        self.assertEqual(cuenta, self.id_cuenta("5.2.08"))

    def test_flujo_web_centralizar_y_f29(self):
        from contabilidad import tributario
        cliente = self.app.test_client()
        self.configurar_otros_impuestos()
        for texto in (CSV_COMPRAS, CSV_VENTAS):
            r = cliente.post("/rcv/importar", data={"periodo": "2026-05",
                                                     "archivo": (io.BytesIO(texto.encode()), "rcv.csv")},
                             content_type="multipart/form-data", follow_redirects=True)
            self.assertIn("documento(s) importado(s)", r.get_data(as_text=True))
        texto = cliente.get("/rcv?periodo=2026-05").get_data(as_text=True)
        self.assertIn("Centralizar compras pendientes", texto)
        self.assertIn("Centralizar ventas pendientes", texto)
        # Centralizar: primero la propuesta (sin guardar), luego guardar
        r = cliente.post("/rcv/documentos", data={"periodo": "2026-05", "centralizar": "VENTA"})
        r = cliente.get(r.headers["Location"])
        texto = r.get_data(as_text=True)
        self.assertIn("aún no se ha guardado", texto)
        self.assertIn("CENTRALIZACIÓN VENTAS MAYO 2026", texto)
        ids = self.conn.execute("SELECT group_concat(id) FROM rcv_documentos WHERE registro = 'VENTA'").fetchone()[0]
        self.assertIn(f'name="rcv_ids" value="{ids}"', texto)
        p = tributario.propuesta_centralizacion(self.conn, self.emp, "2026-05", "VENTA")
        datos = {"tipo": "TRASPASO", "fecha": p["fecha"], "glosa": p["glosa"], "rcv_ids": ids,
                 "cuenta_id": [str(l["cuenta_id"]) for l in p["lineas"]], "glosa_linea": [""] * len(p["lineas"]),
                 "debe": [str(l["debe"] or "") for l in p["lineas"]], "haber": [str(l["haber"] or "") for l in p["lineas"]]}
        self.assertEqual(cliente.post("/asientos/nuevo", data=datos).status_code, 302)
        asiento = self.conn.execute("SELECT DISTINCT asiento_id FROM rcv_documentos WHERE registro = 'VENTA'").fetchall()
        self.assertEqual(len(asiento), 1)
        self.assertEqual(self.saldos_asiento(asiento[0][0]),
                         {"1.1.04": 1_866_000, "4.1.01": -1_600_000, "2.1.04": -266_000})
        # Las compras se centralizan igual
        p = tributario.propuesta_centralizacion(self.conn, self.emp, "2026-05", "COMPRA")
        nucleo_id = nucleo.guardar_asiento(self.conn, self.emp, "TRASPASO", p["fecha"], p["glosa"], p["lineas"])
        tributario.marcar_centralizados(self.conn, self.emp, p["documentos"], nucleo_id)
        # Retenciones del mes y un ajuste de IVA (solo cuentas de IVA: no afecta la comparación)
        self.asiento("EGRESO", "2026-05-15", "Honorarios", ("5.2.03", 500_000, 0),
                     ("2.1.06", 0, 50_000), ("1.1.02", 0, 450_000))
        self.asiento("EGRESO", "2026-05-30", "Sueldos", ("5.2.01", 800_000, 0),
                     ("2.1.05", 0, 12_000), ("1.1.02", 0, 788_000))
        self.asiento("TRASPASO", "2026-05-31", "Ajuste IVA", ("2.1.04", 266_000, 0), ("1.1.09", 0, 266_000))
        f = tributario.calcular_f29(self.conn, self.emp, "2026-05", remanente_anterior=10_000,
                                    utm_anterior=60_000, utm_actual=61_200)
        c = f["codigos"]
        self.assertEqual((c["503"], c["502"], c["509"], c["510"], c["538"]), (2, 285_000, 1, 19_000, 266_000))
        self.assertEqual((c["519"], c["520"], c["524"], c["525"], c["527"], c["528"]),
                         (3, 231_800, 1, 95_000, 1, 19_000))
        self.assertEqual((f["reajuste"], c["504"], c["537"]), (200, 10_200, 318_000))
        self.assertEqual((c["89"], c["77"]), (0, 52_000))
        self.assertEqual((c["563"], c["62"]), (1_600_000, 4_000))
        self.assertEqual((c["151"], c["48"], c["91"]), (50_000, 12_000, 66_000))
        self.assertEqual(f["alertas"], [])  # todo centralizado y cuadrado
        # La página, el PDF y el guardado (el remanente pasa al mes siguiente)
        texto = cliente.get("/f29?periodo=2026-05&remanente_anterior=10.000&utm_anterior=60.000&utm_actual=61.200"
                            ).get_data(as_text=True)
        self.assertIn("52.000", texto)
        self.assertIn("cuadra", texto)
        r = cliente.get("/f29?periodo=2026-05&remanente_anterior=10000&utm_anterior=60000&utm_actual=61200&formato=pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        cliente.post("/f29", data={"periodo": "2026-05", "remanente_anterior": "10000",
                                   "utm_anterior": "60000", "utm_actual": "61200"})
        iniciales = tributario.valores_iniciales_f29(self.conn, self.emp, "2026-06")
        self.assertEqual((iniciales["remanente_anterior"], iniciales["utm_anterior"]), (52_000, 61_200))
        # Borrar el asiento de centralización deja los documentos pendientes otra vez
        nucleo.eliminar_asiento(self.conn, self.emp, nucleo_id)
        pendientes = self.conn.execute(
            "SELECT COUNT(*) FROM rcv_documentos WHERE registro = 'COMPRA' AND asiento_id IS NULL").fetchone()[0]
        self.assertEqual(pendientes, 6)

    def test_resumen_de_boletas_y_comprobantes_de_pago(self):
        from contabilidad import tributario
        self.importar(CSV_VENTAS)
        # Boletas: neto e IVA calculados desde el total
        tributario.registrar_resumen_ventas(self.conn, self.emp, "2026-05", 39, 120, 1_190_000)
        # Vouchers con neto e IVA tomados del resumen del SII
        tributario.registrar_resumen_ventas(self.conn, self.emp, "2026-05", 48, 30, 595_000, neto=500_000, iva=95_000)
        tributario.registrar_resumen_ventas(self.conn, self.emp, "2026-05", 41, 10, 50_000)  # boletas exentas
        boleta = self.conn.execute("SELECT * FROM rcv_documentos WHERE tipo_doc = 39").fetchone()
        self.assertEqual((boleta["neto"], boleta["iva"], boleta["cantidad"], boleta["rut"]),
                         (1_000_000, 190_000, 120, "66.666.666-6"))
        exenta = self.conn.execute("SELECT * FROM rcv_documentos WHERE tipo_doc = 41").fetchone()
        self.assertEqual((exenta["exento"], exenta["neto"], exenta["iva"]), (50_000, 0, 0))
        # Volver a ingresar el mismo tipo lo reemplaza
        tributario.registrar_resumen_ventas(self.conn, self.emp, "2026-05", 39, 125, 1_309_000)
        self.assertEqual(self.conn.execute("SELECT COUNT(*), SUM(cantidad) FROM rcv_documentos WHERE tipo_doc = 39"
                                           ).fetchone()[:], (1, 125))
        for malo in ({"cantidad": 0, "total": 1000}, {"cantidad": 1, "total": 0},
                     {"cantidad": 1, "total": 1000, "neto": 500}, {"cantidad": 1, "total": 1000, "neto": 800, "iva": 100}):
            with self.subTest(malo=malo), self.assertRaises(ErrorValidacion):
                tributario.registrar_resumen_ventas(self.conn, self.emp, "2026-05", 39,
                                                    malo["cantidad"], malo["total"], neto=malo.get("neto"),
                                                    iva=malo.get("iva"))
        # F29
        f = tributario.calcular_f29(self.conn, self.emp, "2026-05")
        c = f["codigos"]
        self.assertEqual((c["110"], c["111"]), (125, 209_000))   # 1.309.000 = 1.100.000 + 209.000
        self.assertEqual((c["758"], c["759"]), (30, 95_000))
        self.assertEqual(c["538"], 285_000 + 209_000 + 95_000 - 19_000)
        self.assertEqual(c["142"], 200_000 + 50_000)             # factura exenta + boletas exentas
        self.assertEqual(c["563"], 1_600_000 + 1_100_000 + 500_000 + 50_000)
        # Centralización: el total de boletas y vouchers va a su cuenta (aquí, Caja)
        tributario.guardar_cuentas_tributarias(self.conn, self.emp, {"ventas_boletas": self.id_cuenta("1.1.01")})
        p = tributario.propuesta_centralizacion(self.conn, self.emp, "2026-05", "VENTA")
        lineas = {nucleo.obtener_cuenta(self.conn, self.emp, l["cuenta_id"])["codigo"]: l["debe"] - l["haber"]
                  for l in p["lineas"]}
        self.assertEqual(lineas["1.1.01"], 1_309_000 + 595_000 + 50_000)  # Caja
        self.assertEqual(lineas["1.1.04"], 1_866_000)                      # Clientes (facturas)
        self.assertEqual(lineas["2.1.04"], -(266_000 + 209_000 + 95_000))  # IVA débito
        self.assertEqual(sum(lineas.values()), 0)
        # Una vez centralizado ya no se puede reemplazar
        asiento = nucleo.guardar_asiento(self.conn, self.emp, "TRASPASO", p["fecha"], p["glosa"], p["lineas"])
        tributario.marcar_centralizados(self.conn, self.emp, p["documentos"], asiento)
        with self.assertRaises(ErrorValidacion):
            tributario.registrar_resumen_ventas(self.conn, self.emp, "2026-05", 39, 1, 1000)

    # Formato real del CSV de resumen de ventas que descarga el SII
    CSV_RESUMEN_SII = ("Tipo Documento;Total Documentos;Monto Exento;Monto Neto;Monto IVA;Monto Total\r\n"
                       "Total Oper. del mes Boleta Electr.(39);17;0;141598;26902;168500\r\n")

    def test_resumen_real_del_sii(self):
        from contabilidad import tributario
        registro, filas = tributario.leer_csv_rcv(self.CSV_RESUMEN_SII.encode("cp1252"))
        self.assertEqual(registro, "RESUMEN")
        self.assertEqual([(f["tipo_doc"], f["cantidad"], f["neto"], f["iva"], f["total"]) for f in filas],
                         [(39, 17, 141_598, 26_902, 168_500)])
        # Con facturas y vouchers: solo se registran boletas y comprobantes de pago (las facturas vienen en el detalle)
        texto = self.CSV_RESUMEN_SII + ("Factura Electronica(33);4;0;1000000;190000;1190000\r\n"
                                        "Comprobante Pago Electronico(48);9;0;100000;19000;119000\r\n"
                                        "Total;30;0;1241598;235902;1477500\r\n")
        self.assertEqual(tributario.importar_rcv(self.conn, self.emp, "2026-09", texto.encode()), ("RESUMEN", 2, 1))
        boleta = self.conn.execute("SELECT neto, iva, cantidad FROM rcv_documentos WHERE tipo_doc = 39").fetchone()
        self.assertEqual(tuple(boleta), (141_598, 26_902, 17))  # montos exactos del SII, sin recalcular
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM rcv_documentos WHERE tipo_doc = 33").fetchone()[0], 0)
        # Un resumen de compras no se acepta
        with self.assertRaises(ErrorValidacion):
            tributario.leer_csv_rcv(b"Tipo Documento;Total Documentos;Monto Exento;Monto Neto;IVA Recuperable;"
                                    b"Monto Total\nFactura(33);1;0;100;19;119\n")

    def test_rut_y_periodo_del_nombre_del_archivo(self):
        cliente = self.app.test_client()
        cliente.post("/empresa", data={"razon_social": "Pyme Ejemplo SpA", "rut": "76086428-5", "ejercicio": "2026"})

        def subir(nombre, periodo="2026-05"):
            r = cliente.post("/rcv/importar", data={"periodo": periodo,
                                                     "archivo": (io.BytesIO(self.CSV_RESUMEN_SII.encode()), nombre)},
                             content_type="multipart/form-data", follow_redirects=True)
            return r.get_data(as_text=True)

        texto = subir("RCV_RESUMEN_VENTA_12345678_202609.csv")  # de otra empresa
        self.assertIn("No se importó", texto)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM rcv_documentos").fetchone()[0], 0)
        texto = subir("RCV_RESUMEN_VENTA_76086428_202609.csv")  # período tomado del nombre
        self.assertIn("se importó en ese período", texto)
        self.assertEqual(self.conn.execute("SELECT periodo FROM rcv_documentos").fetchone()[0], "2026-09")
        subir("resumen.csv", periodo="2026-08")  # sin RUT ni período en el nombre: usa el de la pantalla
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM rcv_documentos WHERE periodo = '2026-08'").fetchone()[0], 1)

    def test_resumen_por_la_web(self):
        cliente = self.app.test_client()
        r = cliente.post("/rcv/resumen", data={"periodo": "2026-05", "tipo_doc": "48", "cantidad": "12",
                                               "total": "238.000"}, follow_redirects=True)
        texto = r.get_data(as_text=True)
        self.assertIn("Resumen de comprobante de pago electrónico de mayo 2026 registrado", texto)
        self.assertIn("(12 doc.)", texto)
        r = cliente.post("/rcv/resumen", data={"periodo": "2026-05", "tipo_doc": "39", "cantidad": "1", "total": "abc"},
                         follow_redirects=True)
        self.assertIn("Monto inválido", r.get_data(as_text=True))

    def test_f29_avisa_diferencias(self):
        from contabilidad import tributario
        self.importar(CSV_VENTAS)
        f = tributario.calcular_f29(self.conn, self.emp, "2026-05")
        texto = " ".join(f["alertas"])
        self.assertIn("sin centralizar", texto)
        self.assertIn("IVA débito fiscal: el RCV indica $266.000 y la contabilidad $0", texto)
        self.assertIn("tasa de PPM es 0", texto)

    def test_eliminar_pendientes_y_paginas(self):
        self.importar()
        cliente = self.app.test_client()
        ids = [r[0] for r in self.conn.execute("SELECT id FROM rcv_documentos LIMIT 2")]
        cliente.post("/rcv/documentos", data={"periodo": "2026-05", "eliminar": "1", "seleccion": ids})
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM rcv_documentos").fetchone()[0], 4)
        for url in ("/rcv", "/rcv?periodo=2026-05", "/tributario/cuentas", "/f29", "/f29?periodo=2026-05",
                    "/f29?periodo=2026-05&formato=xlsx"):
            with self.subTest(url=url):
                self.assertEqual(cliente.get(url).status_code, 200)
        r = cliente.post("/tributario/cuentas", data={"tasa_ppm": "abc"}, follow_redirects=True)
        self.assertIn("tasa de PPM debe ser un número", r.get_data(as_text=True))


class PruebasRestaurar(BaseConDatos):
    def subir(self, cliente, contenido, nombre="respaldo.db"):
        return cliente.post("/respaldos/subir", data={"archivo": (io.BytesIO(contenido), nombre)},
                            content_type="multipart/form-data", follow_redirects=True)

    def test_rechaza_archivo_que_no_es_base_de_datos(self):
        r = self.subir(self.app.test_client(), b"hola, esto no es una base de datos")
        self.assertIn("no es una base de datos SQLite", r.get_data(as_text=True))

    def test_rechaza_sqlite_de_otro_programa(self):
        otra = Path(self.tmp.name) / "otra.db"
        c = sqlite3.connect(otra)
        c.execute("CREATE TABLE clientes (id INTEGER)")
        c.commit()
        c.close()
        r = self.subir(self.app.test_client(), otra.read_bytes())
        self.assertIn("no es un respaldo de este programa", r.get_data(as_text=True))

    def test_restaurar_y_deshacer(self):
        cliente = self.app.test_client()
        self.cargar_ejemplo()
        self.conn.execute("UPDATE empresas SET razon_social = 'Empresa Respaldada'")
        self.conn.commit()
        respaldo = cliente.get("/respaldo").data
        self.conn.execute("DELETE FROM asientos")
        self.conn.execute("UPDATE empresas SET razon_social = 'Empresa Actual'")
        self.conn.commit()

        r = self.subir(cliente, respaldo)
        texto = r.get_data(as_text=True)
        self.assertIn("Confirmar restauración", texto)
        self.assertIn("Empresa Respaldada", texto)
        self.assertIn("Empresa Actual", texto)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0], 0)

        r = cliente.post("/respaldos/restaurar", follow_redirects=True)
        self.assertIn("Respaldo restaurado", r.get_data(as_text=True))
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0], 4)
        self.assertEqual(self.conn.execute("SELECT razon_social FROM empresas").fetchone()[0], "Empresa Respaldada")
        bg = nucleo.balance_general(self.conn, self.emp, "2026-12-31")
        self.assertEqual(bg["total_activo"], bg["total_pasivo_patrimonio"])

        copias = list((self.ruta.parent / "respaldos").glob("antes_de_restaurar_*.db"))
        self.assertEqual(len(copias), 1)
        cliente.post(f"/respaldos/copia/{copias[0].name}")
        cliente.post("/respaldos/restaurar")
        self.assertEqual(self.conn.execute("SELECT razon_social FROM empresas").fetchone()[0], "Empresa Actual")
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM asientos").fetchone()[0], 0)

    def test_restaurar_sin_archivo_pendiente(self):
        r = self.app.test_client().post("/respaldos/restaurar", follow_redirects=True)
        self.assertIn("No hay un archivo pendiente", r.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
