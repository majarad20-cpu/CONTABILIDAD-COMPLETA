"""Exportación de informes a PDF, Excel y CSV.

Cada informe se describe una sola vez como un `Informe` (columnas y filas con
un estilo) y luego se convierte a cualquiera de los tres formatos.

Formato de los libros e informes oficiales (Diario, Mayor, Balance de 8
columnas, Estado de Resultados y Estado de Situación Financiera):
- Encabezado con razón social, RUT, giro, dirección y ciudad, y FOLIO
  correlativo por página en el PDF.
- Tabla con grilla completa y montos con $.
- Declaración del Artículo 100 del Código Tributario y firmas del contador
  general (con imagen de firma, si se cargó) y del representante legal.
- Papel carta.
"""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .nucleo import NOMBRES_TIPO, formato_clp, numero_comprobante

FORMATOS = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
}

NOMBRE_PROGRAMA = "CONTABILIDAD COMPLETA"
ARTICULO_100 = ("Artículo 100 Código Tributario: Declaración jurada de que los datos corresponden "
                "a los registros contables de la empresa.")
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

GRIS_FONDO = "EEEEEE"
GRIS_GRILLA = "9E9E9E"
GRIS_TEXTO = "555555"
ROJO = "B42318"

# Estilos de fila:
#   ""              fila normal
#   seccion         título de sección (fondo gris, abarca toda la fila)
#   glosa           encabezado de asiento en el Libro Diario
#   subtotal        totales intermedios (negrita, línea superior)
#   total           totales finales (negrita, línea superior gruesa)
#   nota            texto informativo en cursiva
NEGRITA = {"seccion", "glosa", "subtotal", "total"}


@dataclass
class Columna:
    titulo: str
    ancho: float            # en caracteres aproximados (Excel) y peso relativo (PDF)
    tipo: str = "texto"     # texto | monto | numero | fecha
    negrita: bool = False   # columna siempre en negrita (p. ej. Saldo del Mayor)


@dataclass
class Fila:
    valores: list
    estilo: str = ""
    nivel: int = 0          # sangría de la columna de texto principal
    abarca: int = None      # la celda en esta posición se extiende hasta la última columna
    derecha: bool = False   # alinear a la derecha la etiqueta de texto (totales)


@dataclass
class Informe:
    titulo: str
    subtitulo: str
    columnas: list
    nombre_archivo: str
    filas: list = field(default_factory=list)
    grupos: list = field(default_factory=list)   # encabezado superior: [(titulo, n_columnas)]
    horizontal: bool = False
    oficial: bool = False   # libros e informes con folio, Art. 100 y firmas de contador y representante
    firmas: list = field(default_factory=list)   # firmas simples (p. ej. comprobantes)
    vacio: str = ""         # cómo se muestra un monto vacío
    notas: list = field(default_factory=list)
    csv_columnas: list = None                    # versión plana opcional para CSV
    csv_filas: list = None

    def agregar(self, *valores, estilo="", nivel=0, abarca=None, derecha=False):
        self.filas.append(Fila(list(valores), estilo, nivel, abarca, derecha))


def fecha_cl(iso):
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}" if iso else ""


def periodo(desde, hasta):
    return f"PERÍODO: {fecha_cl(desde)} a {fecha_cl(hasta)}"


def pesos(valor):
    return f"-${formato_clp(-valor)}" if valor < 0 else f"${formato_clp(valor)}"


def nombre_seguro(texto):
    return re.sub(r"[^a-z0-9]+", "_", texto.lower()).strip("_")[:60] or "empresa"


def _sin_ceros(v):
    return v if v else None


# ==========================================================================
# Definición de cada informe
# ==========================================================================

def informe_libro_diario(d, desde, hasta):
    inf = Informe("LIBRO DIARIO", periodo(desde, hasta),
                  [Columna("Fecha", 11, "fecha"), Columna("As. N°", 7), Columna("Código", 10),
                   Columna("Cuenta / Glosa / Detalle", 46), Columna("Debe", 15, "monto"), Columna("Haber", 15, "monto")],
                  nombre_archivo=f"libro_diario_{desde}_{hasta}", oficial=True)
    inf.csv_columnas = ["Fecha", "Comprobante", "Glosa", "Código", "Cuenta", "Detalle", "Debe", "Haber"]
    inf.csv_filas = []
    mes_actual = None
    for a in d["asientos"]:
        x = a["asiento"]
        if x["fecha"][:7] != mes_actual:
            mes_actual = x["fecha"][:7]
            inf.agregar(f"--- {MESES[int(mes_actual[5:7]) - 1]} {mes_actual[:4]} ---", estilo="seccion", abarca=0)
        marca = " (APERTURA)" if x["apertura"] else " (CIERRE)" if x["cierre"] else ""
        n = numero_comprobante(x["tipo"], x["numero"])
        inf.agregar(x["fecha"], n, f"GLOSA: {x['glosa'].upper()}{marca}", estilo="glosa", abarca=2)
        for l in a["lineas"]:
            detalle = f" · {l['glosa']}" if l["glosa"] else ""
            inf.agregar(None, None, l["codigo"], l["nombre"] + detalle, _sin_ceros(l["debe"]), _sin_ceros(l["haber"]),
                        nivel=2 if l["haber"] else 0)
            inf.csv_filas.append([x["fecha"], n, x["glosa"], l["codigo"], l["nombre"], l["glosa"], l["debe"], l["haber"]])
        inf.agregar(None, None, None, "TOTAL ASIENTO:", sum(l["debe"] for l in a["lineas"]),
                    sum(l["haber"] for l in a["lineas"]), estilo="subtotal", derecha=True)
    inf.agregar(None, None, None, "TOTAL GENERAL LIBRO DIARIO:", d["total_debe"], d["total_haber"],
                estilo="total", derecha=True)
    return inf


def _seccion_mayor(inf, m):
    c = m["cuenta"]
    inf.agregar(f"CUENTA: {c['codigo']} - {c['nombre']}", estilo="seccion", abarca=0)
    if m["saldo_inicial"]:
        inf.agregar(None, None, "SALDO ANTERIOR", None, None, m["saldo_inicial"], estilo="nota")
    for x in m["movimientos"]:
        if x.get("reinicio"):
            inf.agregar(x["fecha"], None, f"Apertura: el saldo anterior ({pesos(x['saldo_anterior'])}) se reinicia",
                        None, None, 0, estilo="nota")
        else:
            inf.agregar(x["fecha"], numero_comprobante(x["tipo"], x["numero"]),
                        (x["glosa"] or x["glosa_asiento"]).upper(),
                        _sin_ceros(x["debe"]), _sin_ceros(x["haber"]), x["saldo"])
    neto = m["saldo_final"] if c["tipo"] in ("ACTIVO", "GASTO") else -m["saldo_final"]
    naturaleza = "SALDO DEUDOR" if neto > 0 else "SALDO ACREEDOR" if neto < 0 else "CUENTA SALDADA"
    inf.agregar(None, None, "TOTALES Y SALDO CUENTA:", m["total_debe"], m["total_haber"],
                (m["saldo_final"], naturaleza), estilo="subtotal", derecha=True)


def _informe_mayor_base(nombre_archivo, subtitulo):
    return Informe("LIBRO MAYOR", subtitulo,
                   [Columna("Fecha", 11, "fecha"), Columna("As. N°", 7), Columna("Glosa / Descripción", 44),
                    Columna("Debe", 14, "monto"), Columna("Haber", 14, "monto"), Columna("Saldo", 15, "monto", negrita=True)],
                   nombre_archivo=nombre_archivo, oficial=True, vacio="-")


def informe_libro_mayor(m, desde, hasta):
    c = m["cuenta"]
    inf = _informe_mayor_base(f"mayor_{c['codigo']}_{desde}_{hasta}", periodo(desde, hasta))
    _seccion_mayor(inf, m)
    return inf


def informe_libro_mayor_completo(datos, desde, hasta):
    inf = _informe_mayor_base(f"libro_mayor_{desde}_{hasta}", periodo(desde, hasta))
    for m in datos["cuentas"]:
        _seccion_mayor(inf, m)
    cuadra = datos["total_debe"] == datos["total_haber"]
    inf.agregar(None, None, "TOTAL GENERAL LIBRO MAYOR:", datos["total_debe"], datos["total_haber"],
                "(CONCORDADO)" if cuadra else "(NO CUADRA)", estilo="total", derecha=True)
    return inf


def informe_balance_ocho(b, desde, hasta):
    cols = ["debitos", "creditos", "deudor", "acreedor", "activo", "pasivo", "perdidas", "ganancias"]
    titulos = ["DEBE", "HABER", "DEUDOR", "ACREEDOR", "ACTIVO", "PASIVO", "PÉRDIDA", "GANANCIA"]
    inf = Informe("BALANCE DE 8 COLUMNAS", periodo(desde, hasta),
                  [Columna("COD.", 8), Columna("CUENTA CONTABLE", 30, negrita=True)]
                  + [Columna(t, 12.5, "monto") for t in titulos],
                  nombre_archivo=f"balance_8_columnas_{desde}_{hasta}",
                  grupos=[("", 2), ("1. SUMAS", 2), ("2. SALDOS", 2), ("3. INVENTARIO", 2), ("4. RESULTADOS", 2)],
                  horizontal=True, oficial=True, vacio="-")
    for f in b["filas"]:
        inf.agregar(f["codigo"], f["nombre"].upper(), f["debitos"], f["creditos"],
                    *[_sin_ceros(f[k]) for k in cols[2:]])
    inf.agregar(None, "SUBTOTALES:", *[b["totales"][k] for k in cols], estilo="subtotal")
    etiqueta = "RESULTADO (UTILIDAD):" if b["resultado"] >= 0 else "RESULTADO (PÉRDIDA):"
    inf.agregar(None, etiqueta, *[_sin_ceros(b["ajuste"][k]) for k in cols], estilo="subtotal")
    inf.agregar(None, "TOTALES GENERALES:", *[b["sumas_iguales"][k] for k in cols], estilo="total")
    return inf


def _hojas(filas):
    """Solo cuentas imputables con saldo, con su código (formato plano)."""
    return [f for f in filas if not f["grupo"] and f["monto"]]


def informe_estado_resultados(er, desde, hasta):
    inf = Informe("ESTADO DE RESULTADOS", periodo(desde, hasta),
                  [Columna("Código", 12), Columna("Cuenta / Detalle", 52), Columna("Resultado", 18, "monto")],
                  nombre_archivo=f"estado_resultados_{desde}_{hasta}", oficial=True)
    inf.agregar("INGRESOS (GANANCIAS)", estilo="seccion", abarca=0)
    for f in _hojas(er["ingresos"]):
        inf.agregar(f["codigo"], f["nombre"], f["monto"])
    inf.agregar(None, "TOTAL INGRESOS", er["total_ingresos"], estilo="subtotal", derecha=True)
    inf.agregar("GASTOS (PÉRDIDAS)", estilo="seccion", abarca=0)
    for f in _hojas(er["gastos"]):
        inf.agregar(f["codigo"], f["nombre"], f["monto"])
    inf.agregar(None, "TOTAL GASTOS", er["total_gastos"], estilo="subtotal", derecha=True)
    inf.agregar(None, "RESULTADO NETO DEL EJERCICIO", er["resultado"], estilo="total", derecha=True)
    return inf


def informe_balance_general(bg, hasta):
    inf = Informe("ESTADO DE SITUACIÓN FINANCIERA", periodo(bg["desde"], hasta),
                  [Columna("Código", 12), Columna("Cuenta / Detalle", 52), Columna("Saldo", 18, "monto")],
                  nombre_archivo=f"estado_situacion_financiera_{hasta}", oficial=True)
    inf.agregar("ACTIVOS", estilo="seccion", abarca=0)
    for f in _hojas(bg["activo"]):
        inf.agregar(f["codigo"], f["nombre"], f["monto"])
    inf.agregar(None, "TOTAL ACTIVOS", bg["total_activo"], estilo="subtotal", derecha=True)
    inf.agregar("PASIVOS", estilo="seccion", abarca=0)
    for f in _hojas(bg["pasivo"]):
        inf.agregar(f["codigo"], f["nombre"], f["monto"])
    inf.agregar(None, "TOTAL PASIVOS", bg["total_pasivo"], estilo="subtotal", derecha=True)
    inf.agregar("PATRIMONIO", estilo="seccion", abarca=0)
    for f in _hojas(bg["patrimonio"]):
        inf.agregar(f["codigo"], f["nombre"], f["monto"])
    if bg["resultado_anterior"]:
        inf.agregar(None, "Resultados anteriores no cerrados", bg["resultado_anterior"])
    inf.agregar(None, "Utilidad / (Pérdida) del Ejercicio", bg["resultado"])
    inf.agregar(None, "TOTAL PATRIMONIO", bg["total_patrimonio"], estilo="subtotal", derecha=True)
    inf.agregar(None, "TOTAL PASIVO + PATRIMONIO", bg["total_pasivo_patrimonio"], estilo="total", derecha=True)
    if bg["total_activo"] != bg["total_pasivo_patrimonio"]:
        inf.notas.append("Atención: el balance no cuadra; revisa los datos.")
    return inf


def informe_comprobante(datos):
    a, lineas = datos["asiento"], datos["lineas"]
    inf = Informe(f"COMPROBANTE DE {a['tipo']} N° {a['numero']}",
                  f"FECHA: {fecha_cl(a['fecha'])} · GLOSA: {a['glosa'].upper()}",
                  [Columna("Código", 11), Columna("Cuenta", 36), Columna("Detalle", 30),
                   Columna("Debe", 15, "monto"), Columna("Haber", 15, "monto")],
                  nombre_archivo=f"comprobante_{numero_comprobante(a['tipo'], a['numero'])}_{a['fecha'][:4]}",
                  firmas=["PREPARADO POR", "REVISADO POR", "APROBADO POR"])
    for l in lineas:
        inf.agregar(l["codigo"], l["nombre"], l["glosa"], _sin_ceros(l["debe"]), _sin_ceros(l["haber"]))
    inf.agregar(None, None, "TOTALES:", sum(l["debe"] for l in lineas), sum(l["haber"] for l in lineas),
                estilo="total", derecha=True)
    if a["apertura"]:
        inf.notas.append("Asiento de apertura: los saldos se reinician desde esta fecha.")
    if a["cierre"]:
        inf.notas.append("Asiento de cierre del ejercicio.")
    return inf


def informe_f29(f):
    """Borrador del F29 (referencial)."""
    from .nucleo import nombre_periodo

    inf = Informe("BORRADOR FORMULARIO 29", f"PERÍODO TRIBUTARIO: {nombre_periodo(f['periodo']).upper()}",
                  [Columna("Código", 9), Columna("Concepto", 52), Columna("Cód.", 7), Columna("Cantidad", 10, "numero"),
                   Columna("Monto", 16, "monto")],
                  nombre_archivo=f"borrador_f29_{f['periodo']}", vacio="")
    c = f["codigos"]
    for linea in f["lineas"]:
        if linea[0] == "titulo":
            inf.agregar(linea[1], estilo="seccion", abarca=0)
            continue
        cod_cantidad, cod_monto, concepto = linea
        if cod_monto == "115":
            inf.agregar(cod_monto, concepto, None, None, f"{c['115']:g} %".replace(".", ","))
        elif cod_cantidad == "total":
            inf.agregar(cod_monto, concepto, None, None, c[cod_monto], estilo="total")
        else:
            cantidad = c[cod_cantidad] if cod_cantidad else None
            inf.agregar(cod_monto, concepto, cod_cantidad, cantidad, c[cod_monto])
    inf.notas.append("Borrador referencial calculado a partir del Registro de Compras y Ventas y de la contabilidad. "
                     "Verifica los valores con la propuesta del SII antes de declarar.")
    inf.notas += f["alertas"]
    return inf


def informe_plan_cuentas(cuentas):
    inf = Informe("PLAN DE CUENTAS", f"{len(cuentas)} CUENTAS",
                  [Columna("Código", 13), Columna("Nombre", 52), Columna("Tipo", 20), Columna("Estado", 10)],
                  nombre_archivo="plan_de_cuentas")
    for c in cuentas:
        inf.agregar(c["codigo"], c["nombre"], NOMBRES_TIPO[c["tipo"]] if c["imputable"] else "Grupo",
                    "" if c["activa"] else "Inactiva",
                    estilo="subtotal" if not c["imputable"] and "." not in c["codigo"] else "",
                    nivel=c["codigo"].count("."))
    return inf


# ==========================================================================
# Datos comunes de encabezado y firmas
# ==========================================================================

def lineas_empresa(empresa):
    return [f"RAZON SOCIAL: {empresa['razon_social']}".upper(),
            f"R.U.T.: {empresa['rut']}",
            f"GIRO: {empresa['giro']}".upper(),
            f"DIRECCION: {empresa['direccion']}".upper(),
            f"CIUDAD: {empresa['comuna']}".upper()]


def _texto_principal(inf):
    """Índice de la columna de texto que recibe sangrías y etiquetas."""
    return max((j for j, c in enumerate(inf.columnas) if c.tipo == "texto"), key=lambda j: inf.columnas[j].ancho)


# ==========================================================================
# CSV
# ==========================================================================

def a_csv(inf):
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";")
    if inf.csv_columnas is not None:
        escritor.writerow(inf.csv_columnas)
        escritor.writerows(inf.csv_filas)
    else:
        escritor.writerow([c.titulo for c in inf.columnas])
        for f in inf.filas:
            valores = f.valores + [None] * (len(inf.columnas) - len(f.valores))
            escritor.writerow(["" if v is None else v[0] if isinstance(v, tuple) else v for v in valores])
    return ("﻿" + salida.getvalue()).encode("utf-8")  # BOM para que Excel reconozca UTF-8


# ==========================================================================
# Excel
# ==========================================================================

def a_excel(inf, empresa, contador):
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "", inf.titulo.title())[:31]
    ws.sheet_view.showGridLines = False
    n = len(inf.columnas)
    ultima = get_column_letter(n)
    principal = _texto_principal(inf)

    linea = Side(style="thin", color=GRIS_GRILLA)
    negro = Side(style="medium", color="000000")
    grilla = Border(left=linea, right=linea, top=linea, bottom=linea)
    fondo = PatternFill("solid", fgColor=GRIS_FONDO)
    formato_monto = '"$"#,##0;[Red]-"$"#,##0'

    fila = 0
    for texto in lineas_empresa(empresa):
        fila += 1
        ws.cell(row=fila, column=1, value=texto).font = Font(bold=True, size=8)
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
    fila += 1
    ws.cell(row=fila, column=1, value=inf.titulo).font = Font(bold=True, size=13)
    ws.cell(row=fila, column=1).alignment = Alignment(horizontal="center")
    ws.merge_cells(f"A{fila}:{ultima}{fila}")
    ws.row_dimensions[fila].height = 20
    fila += 1
    ws.cell(row=fila, column=1, value=inf.subtitulo).font = Font(size=9, color=GRIS_TEXTO)
    ws.cell(row=fila, column=1).alignment = Alignment(horizontal="center")
    ws.merge_cells(f"A{fila}:{ultima}{fila}")

    fila += 2
    primera_encabezado = fila
    if inf.grupos:
        col = 1
        for titulo, span in inf.grupos:
            for j in range(col, col + span):
                celda = ws.cell(row=fila, column=j)
                celda.fill, celda.border = fondo, grilla
            celda = ws.cell(row=fila, column=col, value=titulo or None)
            celda.font = Font(bold=True, size=9)
            if span > 1 and titulo:
                ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + span - 1)
            col += span
        fila += 1
    for j, c in enumerate(inf.columnas, 1):
        celda = ws.cell(row=fila, column=j, value=c.titulo)
        celda.font = Font(bold=True, size=9)
        celda.fill, celda.border = fondo, grilla
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = c.ancho
    if inf.grupos:  # las columnas sin grupo se combinan verticalmente
        col = 1
        for titulo, span in inf.grupos:
            if not titulo:
                for j in range(col, col + span):
                    ws.cell(row=fila - 1, column=j, value=inf.columnas[j - 1].titulo).font = Font(bold=True, size=9)
                    ws.cell(row=fila - 1, column=j).alignment = Alignment(vertical="center")
                    ws.cell(row=fila, column=j, value=None)
                    ws.merge_cells(start_row=fila - 1, start_column=j, end_row=fila, end_column=j)
                break
            col += span
    ws.row_dimensions[fila].height = 18
    fila_titulos = fila
    ws.freeze_panes = f"A{fila + 1}"

    for f in inf.filas:
        fila += 1
        valores = f.valores + [None] * (n - len(f.valores))
        for j, (c, v) in enumerate(zip(inf.columnas, valores), 1):
            sub = None
            if isinstance(v, tuple):
                v, sub = v
            es_dato = f.abarca is None or j - 1 < f.abarca  # no es parte de un texto combinado
            if c.tipo == "fecha" and v and es_dato:
                v = date.fromisoformat(v)
            celda = ws.cell(row=fila, column=j, value=v)
            if v is None and c.tipo == "monto" and inf.vacio and f.abarca is None:
                celda.value = inf.vacio
            negrita = f.estilo in NEGRITA or c.negrita
            celda.font = Font(bold=negrita, italic=f.estilo in ("glosa", "nota"),
                              color=GRIS_TEXTO if f.estilo == "nota" else None, size=9)
            celda.border = grilla
            if f.estilo == "seccion":
                celda.fill = fondo
            if f.estilo in ("subtotal", "total"):
                celda.border = Border(left=linea, right=linea, top=negro, bottom=negro if f.estilo == "total" else linea)
            if c.tipo in ("monto", "numero") and es_dato:
                celda.number_format = formato_monto if c.tipo == "monto" else "#,##0"
                celda.alignment = Alignment(horizontal="right", vertical="top")
                if sub:  # p. ej. «SALDO DEUDOR» bajo el saldo del Mayor
                    etiqueta = ws.cell(row=fila, column=principal + 1)
                    etiqueta.value = f"{etiqueta.value} ({sub})" if etiqueta.value else sub
            elif c.tipo == "fecha" and es_dato:
                celda.number_format = "DD/MM/YYYY"
                celda.alignment = Alignment(horizontal="left", vertical="top")
            else:
                celda.alignment = Alignment(
                    vertical="top", wrap_text=f.abarca is None,
                    horizontal="right" if f.derecha and j - 1 == principal else "left",
                    indent=min(f.nivel, 10) if j - 1 == principal and not f.derecha else 0)
        if f.abarca is not None and f.abarca < n - 1:
            ws.merge_cells(start_row=fila, start_column=f.abarca + 1, end_row=fila, end_column=n)

    for nota in inf.notas:
        fila += 2
        ws.cell(row=fila, column=1, value=nota).font = Font(italic=True, size=8, color=GRIS_TEXTO)
        ws.cell(row=fila, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
        ws.row_dimensions[fila].height = 26

    def bloque_firma(columna_inicio, columna_fin, renglones, fila_linea):
        for j in range(columna_inicio, columna_fin + 1):
            ws.cell(row=fila_linea, column=j).border = Border(top=Side(style="thin", color="000000"))
        for k, texto in enumerate(renglones):
            celda = ws.cell(row=fila_linea + k, column=columna_inicio, value=texto)
            celda.font = Font(bold=True, size=8)
            celda.alignment = Alignment(horizontal="center")
            if columna_fin > columna_inicio:
                ws.merge_cells(start_row=fila_linea + k, start_column=columna_inicio,
                               end_row=fila_linea + k, end_column=columna_fin)

    mitad = max(n // 2, 1)
    if inf.oficial:
        fila += 2
        ws.cell(row=fila, column=1, value=ARTICULO_100).font = Font(bold=True, size=8)
        ws.merge_cells(f"A{fila}:{ultima}{fila}")
        fila += 5
        if contador["firma"]:
            from openpyxl.drawing.image import Image as ImagenExcel

            imagen = ImagenExcel(io.BytesIO(contador["firma"]))
            escala = 55 / imagen.height
            imagen.height, imagen.width = 55, imagen.width * escala
            ws.add_image(imagen, f"A{fila - 3}")
        renglones = ["CONTADOR GENERAL"] + [x for x in (contador["nombre"],
                                                        f"R.U.T: {contador['rut']}" if contador["rut"] else "") if x]
        bloque_firma(1, max(mitad - 1, 1), renglones, fila)
        bloque_firma(min(mitad + 1, n), n, ["REPRESENTANTE LEGAL"]
                     + ([empresa["representante"].upper()] if empresa["representante"] else []), fila)
        fila += len(renglones)
    elif inf.firmas:
        fila += 4
        tramo = max(n // len(inf.firmas), 1)
        for i, firma in enumerate(inf.firmas):
            inicio = 1 + i * tramo
            bloque_firma(inicio, min(inicio + max(tramo - 2, 0), n), [firma], fila)

    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = "landscape" if inf.horizontal else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{primera_encabezado}:{fila_titulos}"
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.oddFooter.left.text = NOMBRE_PROGRAMA
    ws.oddFooter.left.size = 7
    ws.oddFooter.right.text = "PÁGINA &P DE &N"
    ws.oddFooter.right.size = 7
    wb.properties.title = inf.titulo
    wb.properties.creator = empresa["razon_social"]

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ==========================================================================
# PDF
# ==========================================================================

def _color(hex6):
    return colors.HexColor("#" + hex6)


def _ajustar(canv, texto, fuente, tamano, ancho):
    """Acorta el texto con «…» para que quepa en el ancho indicado."""
    if canv.stringWidth(texto, fuente, tamano) <= ancho:
        return texto
    while texto and canv.stringWidth(texto + "…", fuente, tamano) > ancho:
        texto = texto[:-1]
    return texto.rstrip(" ·,") + "…"


def _dividir_estilos(estilos, corte):
    """Reparte los comandos de estilo de una tabla que se divide en la fila
    `corte`: devuelve (estilos de la parte superior, estilos de la cola, con
    las filas renumeradas desde 0)."""
    superior, cola = [], []
    for nombre, (c0, r0), (c1, r1), *resto in estilos:
        if r1 == -1:  # hasta el final de la tabla: aplica a ambas partes
            superior.append((nombre, (c0, r0), (c1, r1), *resto))
            cola.append((nombre, (c0, 0), (c1, -1), *resto))
        elif r1 < corte:
            superior.append((nombre, (c0, r0), (c1, r1), *resto))
        elif r0 >= corte:
            cola.append((nombre, (c0, r0 - corte), (c1, r1 - corte), *resto))
    return superior, cola


def a_pdf(inf, empresa, contador, folio_inicial=None):
    """Genera el PDF. Devuelve (bytes, cantidad de páginas)."""
    tamano = landscape(letter) if inf.horizontal else letter
    ancho_pagina, alto_pagina = tamano
    margen = 1.5 * cm
    salida = io.BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=tamano, leftMargin=margen, rightMargin=margen,
                            topMargin=3.5 * cm, bottomMargin=1.7 * cm,
                            title=f"{inf.titulo.title()} – {empresa['razon_social']}", author=empresa["razon_social"])
    ancho_util = ancho_pagina - 2 * margen
    peso = sum(c.ancho for c in inf.columnas)
    anchos = [ancho_util * c.ancho / peso for c in inf.columnas]
    n = len(inf.columnas)
    principal = _texto_principal(inf)

    cuerpo = 7 if n > 7 else 7.5
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=cuerpo, leading=cuerpo + 1.5)
    estilos_p = {
        "": base,
        "negrita": ParagraphStyle("negrita", parent=base, fontName="Helvetica-Bold"),
        "glosa": ParagraphStyle("glosa", parent=base, fontName="Helvetica-BoldOblique"),
        "nota": ParagraphStyle("nota", parent=base, fontName="Helvetica-Oblique", textColor=_color(GRIS_TEXTO)),
    }
    titulo_col = ParagraphStyle("titulo_col", parent=base, fontName="Helvetica-Bold")

    def estilo_monto(negrita, negativo):
        """Los montos también son párrafos, para alinearse igual que el texto de la fila."""
        return ParagraphStyle(f"monto_{negrita}_{negativo}", parent=base, alignment=TA_RIGHT,
                              fontName="Helvetica-Bold" if negrita else "Helvetica",
                              textColor=_color(ROJO) if negativo else colors.black)

    datos = []
    estilos = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), cuerpo),
        ("TOPPADDING", (0, 0), (-1, -1), 2.3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("GRID", (0, 0), (-1, -1), 0.5, _color(GRIS_GRILLA)),
    ]
    if inf.grupos:
        fila_grupos, col = [], 0
        for titulo, span in inf.grupos:
            fila_grupos += [Paragraph(escape(titulo), titulo_col)] + [""] * (span - 1)
            if span > 1 and titulo:
                estilos.append(("SPAN", (col, 0), (col + span - 1, 0)))
            col += span
        datos.append(fila_grupos)
        col = 0
        for titulo, span in inf.grupos:  # columnas sin grupo: título combinado en las dos filas
            if not titulo:
                for j in range(col, col + span):
                    fila_grupos[j] = Paragraph(escape(inf.columnas[j].titulo), titulo_col)
                    estilos.append(("SPAN", (j, 0), (j, 1)))
            col += span
    datos.append([Paragraph(escape(c.titulo), titulo_col) for c in inf.columnas])
    n_encabezado = len(datos)
    estilos.append(("BACKGROUND", (0, 0), (-1, n_encabezado - 1), _color(GRIS_FONDO)))

    for i, f in enumerate(inf.filas, start=n_encabezado):
        valores = f.valores + [None] * (n - len(f.valores))
        if f.estilo == "glosa":
            clave_p = "glosa"
        elif f.estilo == "nota":
            clave_p = "nota"
        elif f.estilo in NEGRITA:
            clave_p = "negrita"
        else:
            clave_p = ""
        celdas = []
        for j, (c, v) in enumerate(zip(inf.columnas, valores)):
            if f.abarca is not None and j > f.abarca:
                celdas.append("")
                continue
            negrita = c.negrita or f.estilo in NEGRITA
            if c.tipo == "numero" and (f.abarca is None or j < f.abarca):
                celdas.append(Paragraph("" if v is None else formato_clp(v), estilo_monto(negrita, False)))
            elif c.tipo == "monto" and (f.abarca is None or j < f.abarca):
                if isinstance(v, tuple):
                    monto, sub = v
                    texto = f"{escape(pesos(monto))}<br/><font size='5'>{escape(sub)}</font>"
                    celdas.append(Paragraph(texto, estilo_monto(True, monto < 0)))
                elif isinstance(v, str):
                    celdas.append(Paragraph(escape(v), estilo_monto(True, False)))
                else:
                    texto = inf.vacio if v is None else pesos(v)
                    celdas.append(Paragraph(escape(texto), estilo_monto(negrita, isinstance(v, int) and v < 0)))
            elif c.tipo == "fecha" and (f.abarca is None or j < f.abarca):
                celdas.append(Paragraph(fecha_cl(v) if v else "", estilos_p["negrita" if negrita else ""]))
            else:
                clave = "negrita" if c.negrita and clave_p == "" else clave_p
                if f.estilo == "glosa" and j < f.abarca:
                    clave = "negrita"  # fecha y número del asiento: negrita sin cursiva
                estilo = estilos_p[clave]
                ajustes = {}
                if j == principal and f.nivel and not f.derecha:
                    ajustes["leftIndent"] = f.nivel * 7
                if f.derecha and j == principal:
                    ajustes["alignment"] = TA_RIGHT
                if ajustes:
                    estilo = ParagraphStyle(f"{clave}_{i}_{j}", parent=estilo, **ajustes)
                celdas.append(Paragraph(escape("" if v is None else str(v)), estilo))
        datos.append(celdas)
        if f.abarca is not None and f.abarca < n - 1:
            estilos.append(("SPAN", (f.abarca, i), (n - 1, i)))
        if f.estilo in NEGRITA:
            estilos.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))
        if f.estilo == "seccion":
            estilos.append(("BACKGROUND", (0, i), (-1, i), _color(GRIS_FONDO)))
        if f.estilo == "glosa":
            estilos.append(("LINEABOVE", (0, i), (-1, i), 0.9, colors.black))
        if f.estilo == "subtotal":
            estilos.append(("LINEABOVE", (0, i), (-1, i), 0.9, colors.black))
            estilos.append(("LINEBELOW", (0, i), (-1, i), 0.9, colors.black))
        elif f.estilo == "total":
            estilos.append(("LINEABOVE", (0, i), (-1, i), 1.4, colors.black))
            estilos.append(("LINEBELOW", (0, i), (-1, i), 1.4, colors.black))

    for j, c in enumerate(inf.columnas):
        if c.tipo in ("monto", "numero"):
            estilos.append(("ALIGN", (j, n_encabezado), (j, -1), "RIGHT"))
            if c.negrita:
                estilos.append(("FONTNAME", (j, n_encabezado), (j, -1), "Helvetica-Bold"))

    # En los informes oficiales, las últimas filas de la tabla van junto al
    # Artículo 100 y las firmas, para que la última página no quede solo con firmas.
    filas_cola = min(4, len(inf.filas) - 1) if inf.oficial else 0
    if filas_cola > 0:
        corte = len(datos) - filas_cola
        estilos_tabla, estilos_cola = _dividir_estilos(estilos, corte)
        tabla = Table(datos[:corte], colWidths=anchos, repeatRows=n_encabezado)
        tabla.setStyle(TableStyle(estilos_tabla))
        cola = Table(datos[corte:], colWidths=anchos)
        cola.setStyle(TableStyle(estilos_cola))
        final = [cola]
    else:
        tabla = Table(datos, colWidths=anchos, repeatRows=n_encabezado)
        tabla.setStyle(TableStyle(estilos))
        final = []
    elementos = [tabla]

    pie = ParagraphStyle("pie", parent=base, fontSize=7, leading=9, textColor=_color(GRIS_TEXTO))
    for texto in inf.notas:
        final += [Spacer(1, 0.3 * cm), Paragraph(escape(texto), pie)]

    centrado = ParagraphStyle("centrado", parent=base, fontName="Helvetica-Bold", alignment=TA_CENTER,
                              fontSize=7.5, leading=9.5)
    if inf.oficial:
        articulo = Paragraph(escape(ARTICULO_100), ParagraphStyle("art100", parent=centrado, alignment=TA_RIGHT))
        firma_contador = ""
        if contador["firma"]:
            lector = ImageReader(io.BytesIO(contador["firma"]))
            ancho_img, alto_img = lector.getSize()
            alto = 1.5 * cm
            ancho = min(alto * ancho_img / alto_img, 5 * cm)
            firma_contador = Image(io.BytesIO(contador["firma"]), width=ancho, height=ancho * alto_img / ancho_img)
        texto_contador = "CONTADOR GENERAL" + "".join(
            f"<br/>{escape(x)}" for x in (contador["nombre"], f"R.U.T: {contador['rut']}" if contador["rut"] else "") if x)
        texto_representante = "REPRESENTANTE LEGAL" + (
            f"<br/>{escape(empresa['representante'].upper())}" if empresa["representante"] else "")
        bloque = [[firma_contador, "", ""], [Paragraph(texto_contador, centrado), "", Paragraph(texto_representante, centrado)]]
        ancho_firma = 6 * cm
        firmas = Table(bloque, colWidths=[ancho_firma, ancho_util - 2 * ancho_firma - 1 * cm, ancho_firma],
                       rowHeights=[1.8 * cm, None], hAlign="CENTER")
        firmas.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, 0), "CENTER"), ("VALIGN", (0, 0), (-1, 0), "BOTTOM"),
            ("VALIGN", (0, 1), (-1, 1), "TOP"),
            ("LINEABOVE", (0, 1), (0, 1), 0.6, colors.black), ("LINEABOVE", (2, 1), (2, 1), 0.6, colors.black),
        ]))
        elementos.append(KeepTogether(final + [Spacer(1, 0.5 * cm), articulo, Spacer(1, 0.3 * cm), firmas]))
    elif inf.firmas:
        k = len(inf.firmas)
        ancho_firma = min(ancho_util / k - 0.8 * cm, 5.5 * cm)
        hueco = (ancho_util / k - ancho_firma) / 2
        etiquetas = []
        for x in inf.firmas:
            etiquetas += ["", Paragraph(escape(x), centrado), ""]
        firmas = Table([[""] * (3 * k), etiquetas], colWidths=[hueco, ancho_firma, hueco] * k,
                       rowHeights=[1.6 * cm, None])
        firmas.setStyle(TableStyle([("LINEABOVE", (3 * j + 1, 1), (3 * j + 1, 1), 0.6, colors.black)
                                    for j in range(k)]))
        elementos += final + [Spacer(1, 0.8 * cm), KeepTogether(firmas)]
    else:
        elementos += final

    def encabezado(canv, _doc):
        canv.saveState()
        arriba = alto_pagina - 1.2 * cm
        canv.setFont("Helvetica-Bold", 6.5)
        for k, texto in enumerate(lineas_empresa(empresa)):
            canv.drawString(margen, arriba - k * 0.28 * cm, _ajustar(canv, texto, "Helvetica-Bold", 6.5, ancho_util * 0.7))
        canv.setFont("Helvetica-Bold", 11)
        canv.drawCentredString(ancho_pagina / 2, arriba - 1.55 * cm, inf.titulo)
        canv.setFont("Helvetica", 7)
        canv.drawCentredString(ancho_pagina / 2, arriba - 1.9 * cm, _ajustar(canv, inf.subtitulo, "Helvetica", 7, ancho_util))
        canv.setStrokeColor(_color(GRIS_GRILLA))
        canv.setLineWidth(0.5)
        canv.line(margen, arriba - 2.1 * cm, ancho_pagina - margen, arriba - 2.1 * cm)
        canv.line(margen, 1.35 * cm, ancho_pagina - margen, 1.35 * cm)
        canv.setFont("Helvetica", 7)
        canv.drawString(margen, 1.0 * cm, NOMBRE_PROGRAMA)
        canv.restoreState()

    paginas = []

    class CanvasNumerado(rl_canvas.Canvas):
        """Agrega «PÁGINA X DE Y» y el folio de cada página (requiere conocer el total)."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            paginas.append(total)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.setFont("Helvetica", 7)
                self.drawRightString(ancho_pagina - margen, 1.0 * cm, f"PÁGINA {self._pageNumber} DE {total}")
                if folio_inicial is not None:
                    self.setFont("Helvetica-Bold", 8)
                    self.drawRightString(ancho_pagina - margen, alto_pagina - 1.2 * cm,
                                         f"FOLIO: {folio_inicial + self._pageNumber - 1}")
                super().showPage()
            super().save()

    doc.build(elementos, onFirstPage=encabezado, onLaterPages=encabezado, canvasmaker=CanvasNumerado)
    return salida.getvalue(), paginas[0]


def exportar(inf, empresa, contador, formato, folio_inicial=None):
    """Devuelve (contenido en bytes, tipo MIME, nombre de archivo, páginas del PDF)."""
    nombre = f"{nombre_seguro(empresa['razon_social'])}_{inf.nombre_archivo}.{formato}"
    if formato == "pdf":
        contenido, paginas = a_pdf(inf, empresa, contador, folio_inicial)
        return contenido, FORMATOS[formato], nombre, paginas
    if formato == "xlsx":
        return a_excel(inf, empresa, contador), FORMATOS[formato], nombre, 0
    if formato == "csv":
        return a_csv(inf), FORMATOS[formato], nombre, 0
    raise ValueError(f"Formato no soportado: {formato}")
